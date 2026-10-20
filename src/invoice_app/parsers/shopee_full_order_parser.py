"""Concrete parser for native Shopee Full Order List XLSX exports."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO
from zipfile import BadZipFile
import re
import warnings

from openpyxl import load_workbook

from src.invoice_app.parsers.shopee_weekly_statement_parser import (
    _populated_bounds,
    _worksheet_archive_path,
)
from src.invoice_app.services.market_context import (
    MarketContext,
    MarketKey,
    resolve_market_context,
)


FULL_ORDER_SHEET = "orders"
REQUIRED_HEADERS = (
    "Order ID",
    "Order Status",
    "Order Creation Date",
    "Parent SKU Reference No.",
    "Product Name",
    "SKU Reference No.",
    "Variation Name",
    "Quantity",
    "Product Subtotal",
)
OPTIONAL_HEADERS = ("Return / Refund Status", "Country")
_HEADER_SPACE = re.compile(r"\s+")


class ShopeeFullOrderParseError(ValueError):
    """The workbook cannot be accepted as one deterministic Full Order source."""


@dataclass(frozen=True)
class ShopeeFullOrderRow:
    order_id: str
    order_status: str
    order_created_at: datetime
    parent_sku: str
    seller_sku: str
    product_name: str
    variation: str
    quantity: int
    product_subtotal: Decimal
    return_refund_status: str
    source_row_number: int


@dataclass(frozen=True)
class ParsedShopeeFullOrder:
    market: MarketContext
    source_filename: str
    source_hash: str
    rows: tuple[ShopeeFullOrderRow, ...]
    dimension_fallback_used: bool

    @property
    def order_ids(self) -> tuple[str, ...]:
        return tuple(dict.fromkeys(row.order_id for row in self.rows))


def parse_shopee_full_order(
    source: str | Path | bytes | bytearray | BinaryIO,
    *,
    market: MarketContext | MarketKey | str,
    source_filename: str | None = None,
) -> ParsedShopeeFullOrder:
    """Parse only the source facts required by Full Order Sync V1."""

    context = resolve_market_context(market)
    if context.key not in {MarketKey.SHOPEE_MY, MarketKey.SHOPEE_SG}:
        raise ShopeeFullOrderParseError(
            f"{context.display_name} is not supported by Shopee Full Order Sync."
        )
    workbook_bytes, filename = _read_source(source, source_filename)
    source_hash = sha256(workbook_bytes).hexdigest()
    if Path(filename).suffix.casefold() != ".xlsx":
        raise ShopeeFullOrderParseError("Shopee Full Order source must be an .xlsx workbook.")
    try:
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore", message="Workbook contains no default style.*", category=UserWarning
            )
            workbook = load_workbook(
                BytesIO(workbook_bytes), read_only=True, data_only=True, keep_links=False
            )
    except (BadZipFile, OSError, ValueError, KeyError) as error:
        raise ShopeeFullOrderParseError(
            f"Unreadable or corrupt Shopee Full Order workbook: {error}"
        ) from error

    try:
        if FULL_ORDER_SHEET not in workbook.sheetnames:
            raise ShopeeFullOrderParseError(
                f"Shopee Full Order workbook is missing the '{FULL_ORDER_SHEET}' worksheet."
            )
        worksheet = workbook[FULL_ORDER_SHEET]
        rows, fallback_used = _worksheet_rows(workbook_bytes, worksheet)
        header_index, columns = _required_header(rows)
        parsed = tuple(
            _parse_row(row, row_number, columns)
            for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2)
            if any(value not in (None, "") for value in row)
        )
    finally:
        workbook.close()
    if not parsed:
        raise ShopeeFullOrderParseError("Shopee Full Order workbook contains no order rows.")
    return ParsedShopeeFullOrder(
        market=context,
        source_filename=filename,
        source_hash=source_hash,
        rows=parsed,
        dimension_fallback_used=fallback_used,
    )


def _worksheet_rows(workbook_bytes: bytes, worksheet: Any) -> tuple[list[tuple[Any, ...]], bool]:
    rows = list(worksheet.iter_rows(values_only=True))
    if _find_header(rows) is not None:
        return rows, False
    path = _worksheet_archive_path(workbook_bytes, worksheet.title)
    max_row, max_column = _populated_bounds(workbook_bytes, path)
    if max_row <= 0 or max_column <= 0:
        return rows, False
    worksheet.reset_dimensions()
    return (
        list(
            worksheet.iter_rows(
                min_row=1,
                max_row=max_row,
                min_col=1,
                max_col=max_column,
                values_only=True,
            )
        ),
        True,
    )


def _required_header(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]]:
    found = _find_header(rows)
    if found is None:
        raise ShopeeFullOrderParseError(
            "Shopee Full Order workbook is missing required column(s): "
            + ", ".join(REQUIRED_HEADERS)
        )
    return found


def _find_header(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]] | None:
    required = {_normalize_header(header): header for header in REQUIRED_HEADERS}
    for row_index, row in enumerate(rows[:50]):
        positions = {
            _normalize_header(value): column_index
            for column_index, value in enumerate(row)
            if value not in (None, "")
        }
        if required.keys() <= positions.keys():
            result = {
                original: positions[normalized]
                for normalized, original in required.items()
            }
            for header in OPTIONAL_HEADERS:
                normalized = _normalize_header(header)
                if normalized in positions:
                    result[header] = positions[normalized]
            return row_index, result
    return None


def _parse_row(
    row: tuple[Any, ...], row_number: int, columns: dict[str, int]
) -> ShopeeFullOrderRow:
    def value(header: str) -> Any:
        index = columns.get(header)
        return row[index] if index is not None and index < len(row) else None

    order_id = _required_text(value("Order ID"), row_number, "Order ID")
    return ShopeeFullOrderRow(
        order_id=order_id,
        order_status=_required_text(value("Order Status"), row_number, "Order Status"),
        order_created_at=_required_datetime(
            value("Order Creation Date"), row_number, "Order Creation Date"
        ),
        parent_sku=_text(value("Parent SKU Reference No.")),
        seller_sku=_text(value("SKU Reference No.")),
        product_name=_required_text(value("Product Name"), row_number, "Product Name"),
        variation=_text(value("Variation Name")),
        quantity=_required_quantity(value("Quantity"), row_number),
        product_subtotal=_required_money(value("Product Subtotal"), row_number),
        return_refund_status=_text(value("Return / Refund Status")),
        source_row_number=row_number,
    )


def _read_source(
    source: str | Path | bytes | bytearray | BinaryIO,
    source_filename: str | None,
) -> tuple[bytes, str]:
    if isinstance(source, (str, Path)):
        path = Path(source)
        return path.read_bytes(), source_filename or path.name
    if isinstance(source, (bytes, bytearray)):
        return bytes(source), source_filename or "shopee-full-order.xlsx"
    filename = source_filename or Path(str(getattr(source, "name", ""))).name
    if hasattr(source, "getvalue"):
        return bytes(source.getvalue()), filename or "shopee-full-order.xlsx"
    position = source.tell() if hasattr(source, "tell") else None
    data = source.read()
    if position is not None and hasattr(source, "seek"):
        source.seek(position)
    return bytes(data), filename or "shopee-full-order.xlsx"


def _required_datetime(value: Any, row_number: int, field: str) -> datetime:
    if isinstance(value, datetime):
        return value
    text = _text(value)
    for pattern in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            continue
    raise ShopeeFullOrderParseError(
        f"orders row {row_number} has invalid {field}: {text!r}."
    )


def _required_quantity(value: Any, row_number: int) -> int:
    text = _text(value)
    try:
        decimal = Decimal(text)
        quantity = int(decimal)
    except (InvalidOperation, ValueError):
        quantity = 0
    if quantity <= 0 or Decimal(quantity) != decimal:
        raise ShopeeFullOrderParseError(
            f"orders row {row_number} has invalid Quantity: {text!r}."
        )
    return quantity


def _required_money(value: Any, row_number: int) -> Decimal:
    text = _text(value).replace(",", "")
    try:
        amount = Decimal(text)
    except InvalidOperation as error:
        raise ShopeeFullOrderParseError(
            f"orders row {row_number} has invalid Product Subtotal: {text!r}."
        ) from error
    if amount.as_tuple().exponent < -2:
        raise ShopeeFullOrderParseError(
            f"orders row {row_number} Product Subtotal has more than two decimals."
        )
    return amount.quantize(Decimal("0.01"))


def _required_text(value: Any, row_number: int, field: str) -> str:
    text = _text(value)
    if not text:
        raise ShopeeFullOrderParseError(
            f"orders row {row_number} has blank {field}."
        )
    return text


def _normalize_header(value: Any) -> str:
    return _HEADER_SPACE.sub(" ", _text(value)).casefold()


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()
