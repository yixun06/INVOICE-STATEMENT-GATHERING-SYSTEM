"""Excel export for the shared Weekly Billing summary model."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.invoice_app.domain.weekly_billing import WeeklyBillingReport, WeeklyBillingSummary
from src.invoice_app.services.weekly_billing_staging import (
    STAGING_DATA_HEADERS,
    build_staging_data_rows,
)


PRODUCT_SUMMARY_HEADERS = (
    "No.",
    "Item/Barcode",
    "Description",
    "Qty",
    "UOM",
    "Unit Price",
    "Dis%",
    "Disc Amt",
    "Amount",
)
FINANCIAL_SUMMARY_HEADERS = ("Description", "Amount")


def export_weekly_billing_summary(summary: WeeklyBillingSummary) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Product Summary"
    sheet.append(PRODUCT_SUMMARY_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{max(1, len(summary.product_rows) + 1)}"

    for row in summary.product_rows:
        sheet.append(
            (
                row.number,
                row.nav,
                row.product_name,
                row.quantity,
                row.uom,
                float(row.unit_price),
                row.discount_percent,
                float(row.discount_amount),
                float(row.amount),
            )
        )
    for column in (6, 8, 9):
        for cells in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            cells[0].number_format = '"RM" #,##0.00'
    widths = {
        "A": 8,
        "B": 18,
        "C": 58,
        "D": 10,
        "E": 10,
        "F": 16,
        "G": 10,
        "H": 16,
        "I": 16,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_weekly_billing_report(
    report: WeeklyBillingReport,
    *,
    generation_date: date | None = None,
) -> bytes:
    """Export the exact Product/Financial previews from one committed batch."""

    if report.product_summary.period != report.financial_summary.period:
        raise ValueError("Weekly Billing Product and Financial Summary batches differ.")
    if not report.financial_summary.export_ready:
        raise ValueError("Financial Summary is not export-ready.")
    export_date = generation_date or date.today()
    staging_rows = build_staging_data_rows(
        report.product_summary,
        generation_date=export_date,
    )

    workbook = Workbook()
    product_sheet = workbook.active
    product_sheet.title = "Product Summary"
    _write_product_summary(product_sheet, report.product_summary)

    staging_sheet = workbook.create_sheet("Staging Data")
    _write_staging_data(staging_sheet, staging_rows)

    financial_sheet = workbook.create_sheet("Financial Summary")
    financial_sheet.append(FINANCIAL_SUMMARY_HEADERS)
    for cell in financial_sheet[1]:
        cell.font = Font(bold=True)
    financial_sheet.freeze_panes = "A2"
    financial_sheet.column_dimensions["A"].width = 54
    financial_sheet.column_dimensions["B"].width = 18
    for row in report.financial_summary.rows:
        financial_sheet.append((row.native_label, None if row.amount is None else float(row.amount)))
        excel_row = financial_sheet.max_row
        description = financial_sheet.cell(excel_row, 1)
        amount = financial_sheet.cell(excel_row, 2)
        description.alignment = Alignment(
            horizontal="left",
            vertical="center",
            indent=1 if row.parent_source_row_number is not None else 0,
        )
        if row.amount is not None:
            amount.number_format = '"RM" #,##0.00;[Red]-"RM" #,##0.00'
        if row.line_type == "TOTAL":
            _emphasize_financial_row(financial_sheet, excel_row, bold=True, fill="D9EAD3")
        elif row.line_type == "SUBTOTAL":
            _emphasize_financial_row(financial_sheet, excel_row, bold=True, fill="E2F0D9")
        elif row.line_type == "SECTION_HEADER":
            _emphasize_financial_row(financial_sheet, excel_row, bold=True, fill="D9EAF7")
        elif row.line_type == "REFERENCE":
            _emphasize_financial_row(financial_sheet, excel_row, bold=False, fill="FFF2CC")
    financial_sheet.auto_filter.ref = f"A1:B{financial_sheet.max_row}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _write_staging_data(sheet, rows) -> None:
    sheet.append(STAGING_DATA_HEADERS)
    for column_index, cell in enumerate(sheet[1], start=1):
        cell.font = Font(name="Calibri", size=11, bold=True)
        if column_index <= 16:
            cell.alignment = Alignment(horizontal="left")
    for row in rows:
        sheet.append((
            row.your_reference,
            row.posting_date,
            row.sell_to_customer_no,
            row.currency_code,
            row.item_type,
            row.nav,
            row.location_code,
            row.quantity,
            row.unit_of_measure_code,
            float(row.unit_price_rsp_excl_gst),
            row.order_date,
            row.shipment_date,
            row.external_doc_no,
            row.customer_outlet_code,
            row.business_unit_code_erp,
            row.project_code_erp,
            row.transfer_to_code,
            row.customer_remark,
            row.customer,
            row.usoft_code,
            row.usoft_product_description,
            row.plan_date,
            row.am_pm,
            row.secondary_type,
            row.quantity_per_unit_of_measure,
            row.line_discount_percent,
        ))
    for row_number in range(2, sheet.max_row + 1):
        for column in (2, 11, 12, 22):
            sheet.cell(row_number, column).number_format = "yyyy\\-mm\\-dd"
            sheet.cell(row_number, column).alignment = Alignment(horizontal="right")
        for column in (8, 10, 13):
            sheet.cell(row_number, column).alignment = Alignment(horizontal="right")
        item_type = sheet.cell(row_number, 5)
        item_type.font = Font(name="Arial", size=10)
        item_type.alignment = Alignment(horizontal="left", vertical="top")
    sheet.sheet_format.defaultRowHeight = 15.75
    for column, width in {
        "A": 20.28515625,
        "B": 13,
        "C": 13,
        "D": 13,
        "E": 13,
        "F": 13,
        "G": 13,
        "H": 13,
        "I": 20.7109375,
        "J": 13,
        "K": 13,
        "L": 13,
        "M": 13,
        "N": 13,
        "O": 13,
        "P": 13,
        "Q": 16,
        "R": 16.85546875,
        "S": 13,
        "T": 19.85546875,
        "U": 55.140625,
        "V": 13,
        "W": 13,
        "X": 13,
        "Y": 23.42578125,
        "Z": 15,
    }.items():
        sheet.column_dimensions[column].width = width


def _write_product_summary(sheet, summary: WeeklyBillingSummary) -> None:
    sheet.append(PRODUCT_SUMMARY_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{max(1, len(summary.product_rows) + 1)}"
    for row in summary.product_rows:
        sheet.append((
            row.number, row.nav, row.product_name, row.quantity, row.uom,
            float(row.unit_price), row.discount_percent, float(row.discount_amount),
            float(row.amount),
        ))
    for column in (6, 8, 9):
        for cells in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            cells[0].number_format = '"RM" #,##0.00'
    for column, width in {
        "A": 8, "B": 18, "C": 58, "D": 10, "E": 10,
        "F": 16, "G": 10, "H": 16, "I": 16,
    }.items():
        sheet.column_dimensions[column].width = width


def _emphasize_financial_row(sheet, row: int, *, bold: bool, fill: str) -> None:
    for cell in sheet[row]:
        cell.font = Font(bold=bold)
        cell.fill = PatternFill("solid", fgColor=fill)
