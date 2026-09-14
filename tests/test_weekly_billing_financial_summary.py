from __future__ import annotations

from dataclasses import replace
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest

from src.invoice_app.domain.historical_invoice import CanonicalInvoiceItem, CanonicalInvoiceOrder
from src.invoice_app.domain.weekly_billing import BillingPeriod
from src.invoice_app.repositories.google_sheets_historical_invoice_repository import (
    _serialize_item,
    _serialize_order,
)
from src.invoice_app.services.uat2_persistence_schema import (
    INVOICE_ITEMS_HEADERS,
    INVOICE_ITEMS_TAB,
    INVOICE_ORDERS_HEADERS,
    INVOICE_ORDERS_TAB,
    STATEMENT_DATA_HEADERS,
    STATEMENT_DATA_TAB,
    STATEMENT_FINANCIAL_COMPONENT_HEADERS,
    STATEMENT_FINANCIAL_COMPONENTS_TAB,
    STATEMENT_SUMMARY_HEADERS,
    STATEMENT_SUMMARY_TAB,
)
from src.invoice_app.services.weekly_billing import (
    WeeklyBillingError,
    build_weekly_billing_dataset,
    build_weekly_billing_financial_summary,
    build_weekly_billing_report,
)
from src.invoice_app.services.weekly_billing_export import export_weekly_billing_report


PERIOD = BillingPeriod(date(2026, 8, 31), date(2026, 9, 6), "batch-golden", "a" * 64)

_SUMMARY = (
    ("1. Total Revenue", "TOTAL", None, "14767.32"),
    ("Merchandise Subtotal", "SUBTOTAL", 1, "15869.36"),
    ("Original product price", "DETAIL", 2, "17671.49"),
    ("Your Seller product promotion", "DETAIL", 2, "-1722.72"),
    ("Refund Amount", "DETAIL", 2, "-79.41"),
    ("Voucher & Rebates", "SUBTOTAL", 1, "-1102.04"),
    ("Rebate Provided by Shopee", "DETAIL", 6, "0.00"),
    ("Voucher Sponsored by Seller", "DETAIL", 6, "-1102.04"),
    ("Cofund Voucher Sponsored by Seller", "DETAIL", 6, "0.00"),
    ("Coin Cashback Sponsored by Seller", "DETAIL", 6, "0.00"),
    ("Cofund Coin Cashback Sponsored by Seller", "DETAIL", 6, "0.00"),
    ("2. Total Expenses", "TOTAL", None, "-4525.18"),
    ("Shipping Subtotal", "SUBTOTAL", 12, "-61.04"),
    ("Shipping Fee Paid by Buyer (excl. SST)", "DETAIL", 13, "744.90"),
    ("Shipping Fee Discount from 3PL", "DETAIL", 13, "0.00"),
    ("Shipping Rebate From Shopee", "DETAIL", 13, "579.40"),
    ("Seller Paid Shipping Fee SST", "DETAIL", 13, "-4.55"),
    ("Actual Shipping Fee", "DETAIL", 13, "-1375.60"),
    ("Reverse Shipping Fee", "DETAIL", 13, "-4.90"),
    ("Reverse Shipping Fee SST", "DETAIL", 13, "-0.29"),
    ("Saver Programme Shipping Fee Savings", "DETAIL", 13, "0.00"),
    ("Return to Seller Shipping Fee", "DETAIL", 13, "0.00"),
    ("Fees & Charges", "SUBTOTAL", 12, "-4464.14"),
    ("AMS Commission Fee", "DETAIL", 23, "-153.78"),
    ("Commission Fee (incl. SST)", "DETAIL", 23, "-1634.11"),
    ("Service Fee (Incl. SST)", "DETAIL", 23, "-1087.24"),
    ("Saver Programme Fee (Incl. SST)", "DETAIL", 23, "0.00"),
    ("Transaction Fee (Incl. SST)", "DETAIL", 23, "-631.97"),
    ("Ads Escrow Top Up Fee", "DETAIL", 23, "-957.04"),
    ("3. Total Released Amount", "TOTAL", None, "10242.14"),
    ("Other Reference Values", "SECTION_HEADER", None, None),
    ("Shipping Fee Promotion by Seller", "REFERENCE", 31, "-82.40"),
)

_COMPONENTS = {
    "Product Price": "15948.77",
    "Refund Amount": "-79.41",
    "Rebate Provided by Shopee": "0.00",
    "Voucher Sponsored by Seller": "-1102.04",
    "Cofund Voucher Sponsored by Seller": "0.00",
    "Coin Cashback Sponsored by Seller": "0.00",
    "Cofund Coin Cashback Sponsored by Seller": "0.00",
    "Shipping Fee Paid by Buyer (excl. SST)": "744.90",
    "Shipping Fee Charged by Logistic Provider": "-1375.60",
    "Seller Paid Shipping Fee SST": "-4.55",
    "Shipping Rebate From Shopee": "579.40",
    "Reverse Shipping Fee": "-4.90",
    "Reverse Shipping Fee SST": "-0.29",
    "Saver Programme Shipping Fee Savings": "0.00",
    "Return to Seller Fee": "0.00",
    "Commission Fee (incl. SST)": "-1634.11",
    "Service Fee (Incl. SST)": "-1087.24",
    "Transaction Fee (Incl. SST)": "-631.97",
    "AMS Commission Fee": "-153.78",
    "Saver Programme Fee (Incl. SST)": "0.00",
    "Ads Escrow Top Up Fee": "-957.04",
}


def _order() -> CanonicalInvoiceOrder:
    return CanonicalInvoiceOrder(
        platform="Shopee", order_id="ORDER-1",
        first_imported_at=datetime(2026, 9, 7, tzinfo=timezone.utc),
    )


def _item() -> CanonicalInvoiceItem:
    return CanonicalInvoiceItem(
        platform="Shopee", order_id="ORDER-1", item_index=0, seller_sku="SKU-1",
        nav="5000001", product_name="Product One", quantity=1,
        unit_price=Decimal("10.00"), line_subtotal=Decimal("8.00"),
    )


def _statement_row() -> tuple[str, ...]:
    values = {header: "" for header in STATEMENT_DATA_HEADERS}
    values.update({
        "statement_batch_id": PERIOD.statement_batch_id,
        "record_type": "ORDER", "sequence_no": "1", "platform": "Shopee",
        "statement_file_hash": PERIOD.statement_file_hash,
        "statement_period_from": PERIOD.statement_period_from.isoformat(),
        "statement_period_to": PERIOD.statement_period_to.isoformat(),
        "statement_order_count": "1", "validation_status": "PASSED",
        "commit_status": "COMMITTED", "order_id": "ORDER-1",
    })
    return tuple(values[header] for header in STATEMENT_DATA_HEADERS)


def _summary_rows() -> tuple[tuple[str, ...], ...]:
    output = []
    for source_row, (label, line_type, parent, amount) in enumerate(_SUMMARY, start=1):
        values = {header: "" for header in STATEMENT_SUMMARY_HEADERS}
        values.update({
            "statement_batch_id": PERIOD.statement_batch_id,
            "statement_file_hash": PERIOD.statement_file_hash, "platform": "Shopee",
            "statement_period_from": PERIOD.statement_period_from.isoformat(),
            "statement_period_to": PERIOD.statement_period_to.isoformat(),
            "statement_source_sheet": "Summary", "statement_source_row_number": str(source_row),
            "native_label": label, "line_type": line_type,
            "parent_source_row_number": "" if parent is None else str(parent),
            "component_amount": "" if amount is None else amount, "currency": "RM",
            "commit_status": "COMMITTED",
        })
        output.append(tuple(values[header] for header in STATEMENT_SUMMARY_HEADERS))
    return tuple(output)


def _component_rows(*, extra_sku: bool = False, adjustment_component: bool = False) -> tuple[tuple[str, ...], ...]:
    output = []
    values_by_name = dict(_COMPONENTS)
    if extra_sku:
        values_by_name["Product Price SKU duplicate"] = "999.99"
    for name, amount in values_by_name.items():
        values = {header: "" for header in STATEMENT_FINANCIAL_COMPONENT_HEADERS}
        values.update({
            "statement_batch_id": PERIOD.statement_batch_id,
            "statement_file_hash": PERIOD.statement_file_hash,
            "record_type": "SKU" if name.endswith("SKU duplicate") else "ORDER",
            "statement_source_sheet": "Income", "statement_source_row_number": "10",
            "platform": "Shopee", "order_id": "ORDER-1", "component_name": name,
            "component_amount": amount,
            "statement_period_from": PERIOD.statement_period_from.isoformat(),
            "statement_period_to": PERIOD.statement_period_to.isoformat(),
            "commit_status": "COMMITTED",
        })
        output.append(tuple(values[header] for header in STATEMENT_FINANCIAL_COMPONENT_HEADERS))
    if adjustment_component:
        values = {header: "" for header in STATEMENT_FINANCIAL_COMPONENT_HEADERS}
        values.update({
            "statement_batch_id": PERIOD.statement_batch_id,
            "statement_file_hash": PERIOD.statement_file_hash,
            "record_type": "ADJUSTMENT", "statement_source_sheet": "Adjustment",
            "statement_source_row_number": "999", "platform": "Shopee",
            "component_name": "Product Price", "component_amount": "122.20",
            "statement_period_from": PERIOD.statement_period_from.isoformat(),
            "statement_period_to": PERIOD.statement_period_to.isoformat(),
            "commit_status": "COMMITTED",
        })
        output.append(tuple(values[header] for header in STATEMENT_FINANCIAL_COMPONENT_HEADERS))
    return tuple(output)


def _dataset(*, summary_rows=None, component_rows=None):
    return build_weekly_billing_dataset({
        INVOICE_ORDERS_TAB: (INVOICE_ORDERS_HEADERS, _serialize_order(_order())),
        INVOICE_ITEMS_TAB: (INVOICE_ITEMS_HEADERS, _serialize_item(_item())),
        STATEMENT_DATA_TAB: (STATEMENT_DATA_HEADERS, _statement_row()),
        STATEMENT_SUMMARY_TAB: (STATEMENT_SUMMARY_HEADERS, *(summary_rows if summary_rows is not None else _summary_rows())),
        STATEMENT_FINANCIAL_COMPONENTS_TAB: (STATEMENT_FINANCIAL_COMPONENT_HEADERS, *(component_rows if component_rows is not None else _component_rows())),
    })


def test_financial_summary_uses_persisted_ordered_native_rows_and_golden_controls():
    summary = build_weekly_billing_financial_summary(_dataset(), PERIOD)
    assert summary.export_ready is True
    assert summary.currency == "RM"
    assert len(summary.rows) == 32
    assert [row.native_label for row in summary.rows] == [row[0] for row in _SUMMARY]
    assert summary.rows[2].parent_source_row_number == 2
    assert summary.rows[0].line_type == "TOTAL"
    assert summary.rows[1].line_type == "SUBTOTAL"
    assert summary.rows[2].line_type == "DETAIL"
    assert summary.rows[30].line_type == "SECTION_HEADER"
    assert summary.rows[31].line_type == "REFERENCE"
    assert summary.rows[30].amount is None
    assert summary.rows[6].amount == Decimal("0.00")
    assert summary.rows[7].amount == Decimal("-1102.04")
    assert [control.passed for control in summary.controls] == [True, True, True, True]
    assert summary.controls[-1].native_amount == Decimal("10242.14")


def test_non_order_components_are_not_counted_in_financial_controls():
    summary = build_weekly_billing_financial_summary(
        _dataset(component_rows=_component_rows(extra_sku=True)), PERIOD
    )
    assert summary.export_ready is True
    assert summary.controls[0].derived_amount == Decimal("15869.36")


def test_adjustment_component_is_excluded_from_native_financial_controls():
    summary = build_weekly_billing_financial_summary(
        _dataset(component_rows=_component_rows(adjustment_component=True)), PERIOD
    )
    assert summary.export_ready is True
    assert summary.controls[0].derived_amount == Decimal("15869.36")


def test_control_mismatch_is_not_export_ready_and_blocks_combined_report():
    rows = list(_summary_rows())
    index = STATEMENT_SUMMARY_HEADERS.index("component_amount")
    mutated = list(rows[0])
    mutated[index] = "14767.31"
    rows[0] = tuple(mutated)
    dataset = _dataset(summary_rows=tuple(rows))
    summary = build_weekly_billing_financial_summary(dataset, PERIOD)
    assert summary.export_ready is False
    assert "Total Revenue control failed" in summary.validation_failures[0]
    with pytest.raises(WeeklyBillingError, match="not export-ready"):
        build_weekly_billing_report(dataset, PERIOD)


def test_missing_or_incomplete_native_summary_fails_closed():
    with pytest.raises(WeeklyBillingError, match="expected 32 native lines, found 0"):
        build_weekly_billing_financial_summary(_dataset(summary_rows=()), PERIOD)
    with pytest.raises(WeeklyBillingError, match="expected 32 native lines, found 31"):
        build_weekly_billing_financial_summary(_dataset(summary_rows=_summary_rows()[:-1]), PERIOD)


def test_unified_workbook_has_exact_two_sheets_and_numeric_financial_amounts():
    report = build_weekly_billing_report(_dataset(), PERIOD)
    workbook = openpyxl.load_workbook(BytesIO(export_weekly_billing_report(report)), data_only=True)
    assert workbook.sheetnames == ["Product Summary", "Financial Summary"]
    assert workbook["Product Summary"].max_column == 9
    finance = workbook["Financial Summary"]
    assert finance.cell(1, 1).value == "Description"
    assert finance.cell(1, 2).value == "Amount"
    assert finance.cell(2, 2).value == 14767.32
    assert finance.cell(32, 2).value is None
    assert finance.cell(33, 2).value == -82.4
    assert report.product_summary.period == report.financial_summary.period == PERIOD
