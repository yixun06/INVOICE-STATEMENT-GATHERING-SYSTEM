from __future__ import annotations

from dataclasses import replace
from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest

from src.invoice_app.domain.weekly_billing import (
    BillingPeriod,
    FinancialSummaryRow,
    ProductSummaryRow,
    WeeklyBillingFinancialSummary,
    WeeklyBillingReport,
    WeeklyBillingSummary,
)
from src.invoice_app.services.weekly_billing_export import (
    PRODUCT_SUMMARY_HEADERS,
    export_weekly_billing_report,
)
from src.invoice_app.services.weekly_billing_staging import (
    STAGING_DATA_HEADERS,
    StagingDataError,
    _aggregate_staging_rows,
    build_staging_data_rows,
)
from src.invoice_app.services.product_price_master import ProductPriceMasterRecord


PERIOD = BillingPeriod(
    date(2026, 8, 31),
    date(2026, 9, 6),
    "batch-golden",
    "a" * 64,
)
GENERATION_DATE = date(2026, 9, 15)


def _product_master_records() -> tuple[ProductPriceMasterRecord, ...]:
    return (
        ProductPriceMasterRecord(
            seller_sku="SKU-PLACEHOLDER",
            parent_sku="",
            product_name="Placeholder source row",
            variation_name="",
            unit_selling_price=Decimal("42.90"),
            source_row=1,
            nav_code="5000000",
            usoft_product_description="Must not be exported for placeholder NAV",
        ),
        ProductPriceMasterRecord(
            seller_sku="SKU-000123",
            parent_sku="",
            product_name="Identity fallback row",
            variation_name="Original",
            unit_selling_price=Decimal("10.00"),
            source_row=2,
            nav_code="000123",
            usoft_product_description="USOFT Identity Product",
        ),
        ProductPriceMasterRecord(
            seller_sku="SKU-000123-SECOND",
            parent_sku="",
            product_name="Same NAV second listing",
            variation_name="",
            unit_selling_price=Decimal("10.00"),
            source_row=3,
            nav_code="000123",
            usoft_product_description="USOFT Identity Product",
        ),
        ProductPriceMasterRecord(
            seller_sku="SKU-5000004",
            parent_sku="",
            product_name="Product 4",
            variation_name="",
            unit_selling_price=Decimal("1.00"),
            source_row=4,
            nav_code="5000004",
            usoft_product_description="USOFT Product 4",
        ),
    )


def _product_rows() -> tuple[ProductSummaryRow, ...]:
    rows = [
        ProductSummaryRow(
            number=1,
            nav="5000000",
            product_name="Oil | 1000ml",
            uom=None,
            unit_price=Decimal("42.90"),
            quantity=539,
            discount_percent=None,
            discount_amount=Decimal("0.00"),
            amount=Decimal("23123.10"),
            source_item_count=1,
        ),
        ProductSummaryRow(
            number=2,
            nav="5000000",
            product_name="Oil | 1000ml",
            uom=None,
            unit_price=Decimal("43.90"),
            quantity=1,
            discount_percent=None,
            discount_amount=Decimal("0.00"),
            amount=Decimal("43.90"),
            source_item_count=1,
        ),
        ProductSummaryRow(
            number=3,
            nav="000123",
            product_name="Identity fallback row | Original",
            uom=None,
            unit_price=Decimal("10.00"),
            quantity=1,
            discount_percent=None,
            discount_amount=Decimal("0.00"),
            amount=Decimal("10.00"),
            source_item_count=1,
        ),
    ]
    rows.extend(
        ProductSummaryRow(
            number=number,
            nav=str(5000000 + number),
            product_name=f"Product {number}",
            uom=None,
            unit_price=Decimal("1.00"),
            quantity=1,
            discount_percent=None,
            discount_amount=Decimal("0.00"),
            amount=Decimal("1.00"),
            source_item_count=1,
        )
        for number in range(4, 178)
    )
    return tuple(rows)


def _summary() -> WeeklyBillingSummary:
    product_rows = _product_rows()
    return WeeklyBillingSummary(
        period=PERIOD,
        order_count=296,
        invoice_item_count=442,
        source_items=(),
        product_rows=product_rows,
        total_quantity=sum(row.quantity for row in product_rows),
        total_standard_amount=sum(
            (row.unit_price * row.quantity for row in product_rows),
            Decimal("0.00"),
        ),
        total_discount_amount=Decimal("0.00"),
        total_amount=sum((row.amount for row in product_rows), Decimal("0.00")),
        normal_amount_total=Decimal("0.00"),
        promotion_amount_total=Decimal("0.00"),
        promotion_group_count=0,
        same_price_promotion_count=0,
        mixed_price_promotion_count=0,
    )


def _report() -> WeeklyBillingReport:
    financial = WeeklyBillingFinancialSummary(
        period=PERIOD,
        currency="MYR",
        rows=(
            FinancialSummaryRow(
                statement_source_row_number=1,
                native_label="1. Total Revenue",
                line_type="TOTAL",
                parent_source_row_number=None,
                amount=Decimal("14767.32"),
                currency="MYR",
            ),
        ),
        controls=(),
        export_ready=True,
        validation_failures=(),
    )
    return WeeklyBillingReport(_summary(), financial)


def test_staging_mapper_is_one_to_one_and_conserves_golden_quantity():
    summary = _summary()

    rows = build_staging_data_rows(
        summary,
        generation_date=GENERATION_DATE,
        product_master_records=_product_master_records(),
    )

    assert len(summary.product_rows) == len(rows) == 177
    assert summary.total_quantity == sum(row.quantity for row in rows) == 715
    assert {row.your_reference for row in rows} == {"SP202608310906"}
    assert {row.posting_date for row in rows} == {GENERATION_DATE}
    assert {row.order_date for row in rows} == {GENERATION_DATE}
    assert {row.shipment_date for row in rows} == {GENERATION_DATE}
    assert (rows[0].nav, rows[0].quantity, rows[0].unit_price_rsp_excl_gst) == (
        summary.product_rows[0].nav,
        summary.product_rows[0].quantity,
        summary.product_rows[0].unit_price,
    )
    assert [(row.nav, row.unit_price_rsp_excl_gst) for row in rows[:2]] == [
        ("5000000", Decimal("42.90")),
        ("5000000", Decimal("43.90")),
    ]
    assert rows[2].nav == "000123"
    assert rows[0].usoft_product_description == "N/A"
    assert rows[1].usoft_product_description == "N/A"
    assert rows[2].usoft_product_description == "USOFT Identity Product"
    assert rows[3].usoft_product_description == "USOFT Product 4"


def test_staging_aggregates_same_nav_and_exact_decimal_price_in_first_occurrence_order():
    source = _summary().product_rows[2]
    product_rows = (
        replace(source, number=1, nav="NAV-A", unit_price=Decimal("12.90"), quantity=9),
        replace(source, number=2, nav="NAV-B", unit_price=Decimal("5.00"), quantity=3),
        replace(source, number=3, nav="NAV-A", unit_price=Decimal("12.90"), quantity=2),
        replace(source, number=4, nav="NAV-A", unit_price=Decimal("12.900"), quantity=4),
    )
    summary = replace(_summary(), product_rows=product_rows, total_quantity=18)

    rows = build_staging_data_rows(summary, generation_date=GENERATION_DATE)

    assert [(row.nav, row.unit_price_rsp_excl_gst, row.quantity) for row in rows] == [
        ("NAV-A", Decimal("12.90"), 15),
        ("NAV-B", Decimal("5.00"), 3),
    ]


def test_staging_keeps_same_nav_in_distinct_exact_price_buckets():
    source = _summary().product_rows[2]
    product_rows = tuple(
        replace(source, number=index, nav="3000573", unit_price=price, quantity=quantity)
        for index, price, quantity in (
            (1, Decimal("39.90"), 1),
            (2, Decimal("36.00"), 11),
            (3, Decimal("7.20"), 7),
        )
    )
    summary = replace(_summary(), product_rows=product_rows, total_quantity=19)

    rows = build_staging_data_rows(summary, generation_date=GENERATION_DATE)

    assert [(row.unit_price_rsp_excl_gst, row.quantity) for row in rows] == [
        (Decimal("39.90"), 1),
        (Decimal("36.00"), 11),
        (Decimal("7.20"), 7),
    ]


def test_staging_does_not_merge_different_nav_or_unproven_placeholder_rows():
    source = _summary().product_rows[2]
    product_rows = (
        replace(source, number=1, nav="NAV-A", unit_price=Decimal("12.90"), quantity=2),
        replace(source, number=2, nav="NAV-B", unit_price=Decimal("12.90"), quantity=3),
        replace(source, number=3, nav="5000000", unit_price=Decimal("12.90"), quantity=4),
        replace(source, number=4, nav="5000000", unit_price=Decimal("12.90"), quantity=5),
    )
    summary = replace(_summary(), product_rows=product_rows, total_quantity=14)

    rows = build_staging_data_rows(summary, generation_date=GENERATION_DATE)

    assert [(row.nav, row.quantity) for row in rows] == [
        ("NAV-A", 2), ("NAV-B", 3), ("5000000", 4), ("5000000", 5)
    ]


def test_staging_aggregation_rejects_conflicting_generated_metadata():
    source = _summary().product_rows[2]
    rows = build_staging_data_rows(
        replace(_summary(), product_rows=(source,), total_quantity=source.quantity),
        generation_date=GENERATION_DATE,
    )

    with pytest.raises(StagingDataError, match="NAV 000123.*customer"):
        _aggregate_staging_rows((rows[0], replace(rows[0], customer="CONFLICT")))


def test_export_aggregates_only_staging_and_preserves_product_and_financial_sheets():
    source = _summary().product_rows[2]
    product_rows = (
        replace(source, number=1, nav="4007457", unit_price=Decimal("12.90"), quantity=9),
        replace(source, number=2, nav="4007457", unit_price=Decimal("12.90"), quantity=2),
    )
    summary = replace(_summary(), product_rows=product_rows, total_quantity=11)
    report = WeeklyBillingReport(summary, _report().financial_summary)

    workbook = openpyxl.load_workbook(
        BytesIO(export_weekly_billing_report(
            report,
            generation_date=GENERATION_DATE,
            product_master_records=(),
        )),
        data_only=True,
    )

    assert [tuple(cell.value for cell in row) for row in workbook["Product Summary"].iter_rows(min_row=2)] == [
        (1, "4007457", "Identity fallback row | Original", 9, None, 12.9, None, 0, 10),
        (2, "4007457", "Identity fallback row | Original", 2, None, 12.9, None, 0, 10),
    ]
    assert workbook["Staging Data"].max_row == 2
    assert workbook["Staging Data"].cell(2, 6).value == "4007457"
    assert workbook["Staging Data"].cell(2, 8).value == 11
    assert [tuple(cell.value for cell in row) for row in workbook["Financial Summary"].iter_rows(min_row=2)] == [
        ("1. Total Revenue", 14767.32),
    ]


def test_staging_mapper_applies_exact_erp_constants_and_blank_outlet():
    row = build_staging_data_rows(
        _summary(),
        generation_date=GENERATION_DATE,
        product_master_records=_product_master_records(),
    )[0]

    assert (
        row.sell_to_customer_no,
        row.currency_code,
        row.item_type,
        row.location_code,
        row.unit_of_measure_code,
        row.customer_outlet_code,
        row.business_unit_code_erp,
        row.project_code_erp,
        row.external_doc_no,
        row.transfer_to_code,
        row.customer,
        row.usoft_code,
        row.plan_date,
    ) == (
        "HC001543",
        "MYR",
        "Item",
        "JH02",
        "EA",
        None,
        "RETAIL",
        "JH02",
        "Shopee",
        None,
        "RETAIL",
        "5000000",
        GENERATION_DATE,
    )


def test_staging_mapper_fails_when_quantity_control_differs():
    summary = _summary()

    with pytest.raises(StagingDataError, match="Quantity total control"):
        build_staging_data_rows(
            replace(summary, total_quantity=714),
            generation_date=GENERATION_DATE,
            product_master_records=_product_master_records(),
        )


def test_staging_mapper_rejects_conflicting_descriptions_for_one_nav():
    records = (
        ProductPriceMasterRecord(
            seller_sku="SKU-A", parent_sku="", product_name="A", variation_name="",
            unit_selling_price=Decimal("1.00"), source_row=10, nav_code="000123",
            usoft_product_description="Description A",
        ),
        ProductPriceMasterRecord(
            seller_sku="SKU-B", parent_sku="", product_name="B", variation_name="",
            unit_selling_price=Decimal("1.00"), source_row=11, nav_code="000123",
            usoft_product_description="Description B",
        ),
    )

    with pytest.raises(StagingDataError, match="conflicting USOFT product descriptions"):
        build_staging_data_rows(
            _summary(),
            generation_date=GENERATION_DATE,
            product_master_records=records,
        )


def test_staging_mapper_uses_approved_row_106_description_for_nav_4001971():
    summary = replace(
        _summary(),
        product_rows=(
            replace(_summary().product_rows[2], nav="4001971"),
        ),
        total_quantity=1,
    )
    records = (
        ProductPriceMasterRecord(
            seller_sku="SKU-ROW-1016", parent_sku="", product_name="Pumpkin", variation_name="",
            unit_selling_price=Decimal("10.00"), source_row=1016, nav_code="4001971",
            usoft_product_description="SN ORG Pumpkin Seed Cube BTL 150g",
        ),
        ProductPriceMasterRecord(
            seller_sku="", parent_sku="PARENT-ROW-106", product_name="Pumpkin", variation_name="",
            unit_selling_price=Decimal("10.00"), source_row=106, nav_code="4001971",
            usoft_product_description="SN Org Pumpkin Seed Cube BTL 150g",
        ),
    )

    rows = build_staging_data_rows(
        summary,
        generation_date=GENERATION_DATE,
        product_master_records=records,
    )

    assert rows[0].usoft_product_description == "SN Org Pumpkin Seed Cube BTL 150g"


def test_unified_workbook_writes_exact_staging_contract_and_types():
    report = _report()
    workbook = openpyxl.load_workbook(
        BytesIO(
            export_weekly_billing_report(
                report,
                generation_date=GENERATION_DATE,
                product_master_records=_product_master_records(),
            )
        ),
        data_only=True,
    )

    assert workbook.sheetnames == ["Product Summary", "Staging Data", "Financial Summary"]
    product = workbook["Product Summary"]
    staging = workbook["Staging Data"]
    financial = workbook["Financial Summary"]
    assert tuple(cell.value for cell in product[1]) == PRODUCT_SUMMARY_HEADERS
    assert tuple(cell.value for cell in staging[1]) == (
        "Your Reference",
        "Posting Date",
        "Sell-To Customer No.",
        "Currency Code",
        "Type",
        "No.",
        "Location Code",
        "Quantity",
        "Unit of Measure Code",
        "Unit Price [RSP] Excl.GST",
        "Order Date",
        "Shipment Date",
        "External Doc No.",
        "Customer Outlet Code",
        "Business Unit Code ERP",
        "Project Code ERP",
        "Transfer-to Code",
        "Customer ",
        "Customer",
        "USOFT - USOFT CODE",
        "USOFT product description",
        "Plan Date",
        "AM/PM",
        "TYPE",
        "Qty. per Unit of Measure",
        "Line Discount %",
    )
    assert staging.max_row == product.max_row == 178
    assert staging.max_column == 26
    assert staging.freeze_panes is None
    assert staging.auto_filter.ref is None
    expected_widths = {
        "A": 27.28515625,
        "B": 22.28515625,
        "H": 18.85546875,
        "I": 20.7109375,
        "J": 23.140625,
        "K": 26.5703125,
        "M": 15.7109375,
        "Q": 16.0,
        "R": 16.85546875,
        "T": 19.85546875,
        "U": 66.7109375,
        "Y": 23.42578125,
        "Z": 15.0,
    }
    for column, width in expected_widths.items():
        assert staging.column_dimensions[column].width == pytest.approx(width)
    assert {
        column
        for column in expected_widths
        if staging.column_dimensions[column].bestFit
    } == {"A", "B", "H", "I", "J", "K", "M", "Q", "R", "T", "Y", "Z"}
    assert staging.cell(1, 1).font.name == "Calibri"
    assert staging.cell(1, 1).font.sz == 11
    assert staging.cell(1, 1).font.bold is True
    assert staging.cell(1, 1).font.color.rgb == "FF000000"
    assert staging.cell(1, 1).alignment.horizontal == "left"
    assert staging.cell(1, 17).alignment.horizontal is None
    assert staging.cell(2, 1).value == "SP202608310906"
    assert staging.cell(2, 2).value.date() == GENERATION_DATE
    assert staging.cell(2, 11).value.date() == GENERATION_DATE
    assert staging.cell(2, 12).value.date() == GENERATION_DATE
    assert staging.cell(2, 6).value == "5000000"
    assert staging.cell(3, 6).value == "5000000"
    assert staging.cell(4, 6).value == "000123"
    assert staging.cell(2, 8).value == 539
    assert staging.cell(2, 10).value == 42.9
    assert staging.cell(3, 10).value == 43.9
    assert staging.cell(2, 5).value == "Item"
    assert staging.cell(2, 13).value == "Shopee"
    assert staging.cell(2, 14).value is None
    assert staging.cell(2, 17).value is None
    assert staging.cell(2, 18).value is None
    assert staging.cell(2, 19).value == "RETAIL"
    assert staging.cell(2, 20).value == "5000000"
    assert staging.cell(2, 21).value == "N/A"
    assert staging.cell(4, 21).value == "USOFT Identity Product"
    assert staging.cell(2, 22).value.date() == GENERATION_DATE
    assert all(staging.cell(2, column).value is None for column in range(23, 27))
    assert staging.cell(2, 2).data_type == "d"
    assert staging.cell(2, 6).data_type == "s"
    assert staging.cell(2, 8).data_type == "n"
    assert staging.cell(2, 10).data_type == "n"
    assert staging.cell(2, 2).number_format == "yyyy\\-mm\\-dd"
    assert staging.cell(2, 6).number_format == "General"
    assert staging.cell(2, 8).number_format == "General"
    assert staging.cell(2, 10).number_format == "General"
    assert staging.cell(2, 22).number_format == "yyyy\\-mm\\-dd"
    assert staging.cell(2, 5).font.name == "Arial"
    assert staging.cell(2, 5).font.sz == 10
    assert staging.cell(2, 5).font.color.rgb == "FF000000"
    assert staging.cell(2, 1).font.color.rgb == "FF000000"
    assert staging.cell(2, 5).alignment.horizontal == "left"
    assert staging.cell(2, 5).alignment.vertical == "top"
    assert sum(staging.cell(row, 8).value for row in range(2, 179)) == 715
    assert product.cell(2, 2).value == "5000000"
    assert product.cell(2, 6).value == 42.9
    assert financial.cell(2, 2).value == 14767.32
