from __future__ import annotations

from dataclasses import replace
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from pypdf import PdfReader

from src.invoice_app.domain.historical_invoice import (
    CanonicalInvoiceItem,
    CanonicalInvoiceOrder,
)
from src.invoice_app.repositories.google_sheets_historical_invoice_repository import (
    _serialize_item,
    _serialize_order,
)
from src.invoice_app.services.market_context import SHOPEE_SG
from src.invoice_app.services.uat2_data_settings import UAT2DataSettings
from src.invoice_app.services.uat2_persistence_schema import (
    INVOICE_ITEMS_HEADERS,
    INVOICE_ITEMS_TAB,
    INVOICE_ORDERS_HEADERS,
    INVOICE_ORDERS_TAB,
    MONTHLY_STATEMENT_DATA_HEADERS,
    MONTHLY_STATEMENT_DATA_TAB,
    MONTHLY_STATEMENT_FINANCIAL_COMPONENT_HEADERS,
    MONTHLY_STATEMENT_FINANCIAL_COMPONENTS_TAB,
    MONTHLY_STATEMENT_SUMMARY_HEADERS,
    MONTHLY_STATEMENT_SUMMARY_TAB,
)


PERIOD_FROM = date(2026, 8, 1)
PERIOD_TO = date(2026, 8, 31)
BATCH_ID = "monthly-august"
FILE_HASH = "m" * 64


def _order(order_id: str = "ORDER-1") -> CanonicalInvoiceOrder:
    return CanonicalInvoiceOrder(
        platform="Shopee",
        order_id=order_id,
        first_imported_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
    )


def _item(order_id: str = "ORDER-1") -> CanonicalInvoiceItem:
    return CanonicalInvoiceItem(
        platform="Shopee",
        order_id=order_id,
        item_index=0,
        seller_sku="9555208000001",
        resolved_seller_sku="9555208000001",
        nav="5000001",
        product_name="Monthly Product",
        quantity=1,
        unit_price=Decimal("10.00"),
        line_subtotal=Decimal("8.00"),
    )


def _monthly_row(
    *,
    batch_id: str = BATCH_ID,
    file_hash: str = FILE_HASH,
    period_from: date = PERIOD_FROM,
    period_to: date = PERIOD_TO,
    commit_status: str = "COMMITTED",
    validation_status: str = "PASSED",
    record_type: str = "ORDER",
    order_id: str = "ORDER-1",
    order_count: str = "1",
    sku_count: str = "1",
) -> tuple[str, ...]:
    values = {header: "" for header in MONTHLY_STATEMENT_DATA_HEADERS}
    values.update(
        {
            "statement_batch_id": batch_id,
            "record_type": record_type,
            "sequence_no": "1" if record_type == "ORDER" else "2",
            "platform": "Shopee",
            "statement_source_filename": "Aug 2026.xlsx",
            "statement_file_hash": file_hash,
            "statement_period_from": period_from.isoformat(),
            "statement_period_to": period_to.isoformat(),
            "statement_order_count": order_count,
            "statement_sku_count": sku_count,
            "validation_status": validation_status,
            "commit_status": commit_status,
            "order_id": order_id if record_type == "ORDER" else "",
        }
    )
    return tuple(values[header] for header in MONTHLY_STATEMENT_DATA_HEADERS)


def _summary_rows() -> tuple[tuple[str, ...], ...]:
    labels = {
        1: ("1. Total Revenue", "TOTAL", "8.00"),
        2: ("Merchandise Subtotal", "SUBTOTAL", "8.00"),
        3: ("2. Total Expenses", "TOTAL", "0.00"),
        4: ("3. Total Released Amount", "TOTAL", "8.00"),
    }
    rows = []
    for source_row_number in range(1, 33):
        label, line_type, amount = labels.get(
            source_row_number,
            (f"Monthly native line {source_row_number}", "DETAIL", "0.00"),
        )
        values = {header: "" for header in MONTHLY_STATEMENT_SUMMARY_HEADERS}
        values.update(
            {
                "statement_batch_id": BATCH_ID,
                "statement_file_hash": FILE_HASH,
                "platform": "Shopee",
                "statement_period_from": PERIOD_FROM.isoformat(),
                "statement_period_to": PERIOD_TO.isoformat(),
                "statement_source_sheet": "Summary",
                "statement_source_row_number": str(source_row_number),
                "native_label": label,
                "line_type": line_type,
                "component_amount": amount,
                "currency": "RM",
                "commit_status": "COMMITTED",
            }
        )
        rows.append(tuple(values[header] for header in MONTHLY_STATEMENT_SUMMARY_HEADERS))
    return tuple(rows)


def _component_rows() -> tuple[tuple[str, ...], ...]:
    component_values = {
        "Product Price": "8.00",
        "Refund Amount": "0.00",
        "Rebate Provided by Shopee": "0.00",
        "Voucher Sponsored by Seller": "0.00",
        "Cofund Voucher Sponsored by Seller": "0.00",
        "Coin Cashback Sponsored by Seller": "0.00",
        "Cofund Coin Cashback Sponsored by Seller": "0.00",
        "Shipping Fee Paid by Buyer (excl. SST)": "0.00",
        "Shipping Fee Charged by Logistic Provider": "0.00",
        "Seller Paid Shipping Fee SST": "0.00",
        "Shipping Rebate From Shopee": "0.00",
        "Reverse Shipping Fee": "0.00",
        "Reverse Shipping Fee SST": "0.00",
        "Saver Programme Shipping Fee Savings": "0.00",
        "Return to Seller Fee": "0.00",
        "Commission Fee (incl. SST)": "0.00",
        "Service Fee (Incl. SST)": "0.00",
        "Transaction Fee (Incl. SST)": "0.00",
        "AMS Commission Fee": "0.00",
        "Saver Programme Fee (Incl. SST)": "0.00",
        "Ads Escrow Top Up Fee": "0.00",
    }
    rows = []
    for sequence_no, (name, amount) in enumerate(component_values.items(), start=1):
        values = {header: "" for header in MONTHLY_STATEMENT_FINANCIAL_COMPONENT_HEADERS}
        values.update(
            {
                "statement_batch_id": BATCH_ID,
                "statement_file_hash": FILE_HASH,
                "record_type": "ORDER",
                "statement_source_sheet": "Income",
                "statement_source_row_number": "2",
                "sequence_no": str(sequence_no),
                "platform": "Shopee",
                "order_id": "ORDER-1",
                "component_name": name,
                "component_amount": amount,
                "statement_period_from": PERIOD_FROM.isoformat(),
                "statement_period_to": PERIOD_TO.isoformat(),
                "commit_status": "COMMITTED",
            }
        )
        rows.append(
            tuple(values[header] for header in MONTHLY_STATEMENT_FINANCIAL_COMPONENT_HEADERS)
        )
    return tuple(rows)


def _tabs(
    monthly_rows: tuple[tuple[str, ...], ...] | None = None,
    *,
    orders: tuple[CanonicalInvoiceOrder, ...] = (_order(),),
    items: tuple[CanonicalInvoiceItem, ...] = (_item(),),
) -> dict[str, tuple[tuple[str, ...], ...]]:
    return {
        INVOICE_ORDERS_TAB: (
            INVOICE_ORDERS_HEADERS,
            *(_serialize_order(order) for order in orders),
        ),
        INVOICE_ITEMS_TAB: (
            INVOICE_ITEMS_HEADERS,
            *(_serialize_item(item) for item in items),
        ),
        MONTHLY_STATEMENT_DATA_TAB: (
            MONTHLY_STATEMENT_DATA_HEADERS,
            *(monthly_rows or (_monthly_row(), _monthly_row(record_type="SKU"))),
        ),
        MONTHLY_STATEMENT_FINANCIAL_COMPONENTS_TAB: (
            MONTHLY_STATEMENT_FINANCIAL_COMPONENT_HEADERS,
            *_component_rows(),
        ),
        MONTHLY_STATEMENT_SUMMARY_TAB: (
            MONTHLY_STATEMENT_SUMMARY_HEADERS,
            *_summary_rows(),
        ),
    }


class _Gateway:
    def __init__(self, tabs):
        self.tabs = tabs
        self.requests: list[tuple[str, tuple[str, ...]]] = []

    def read_tabs(self, spreadsheet_id: str, tabs: tuple[str, ...]):
        self.requests.append((spreadsheet_id, tuple(tabs)))
        return {tab: self.tabs[tab] for tab in tabs}


def test_monthly_reader_requests_exactly_the_five_monthly_source_tabs():
    from src.invoice_app.services.monthly_billing import GoogleSheetsMonthlyBillingReader

    gateway = _Gateway(_tabs())
    dataset = GoogleSheetsMonthlyBillingReader(
        spreadsheet_id="monthly-sheet", gateway=gateway
    ).load_dataset()

    assert gateway.requests == [
        (
            "monthly-sheet",
            (
                INVOICE_ORDERS_TAB,
                INVOICE_ITEMS_TAB,
                MONTHLY_STATEMENT_DATA_TAB,
                MONTHLY_STATEMENT_FINANCIAL_COMPONENTS_TAB,
                MONTHLY_STATEMENT_SUMMARY_TAB,
            ),
        )
    ]
    assert dataset.periods[0].statement_batch_id == BATCH_ID


def test_monthly_dataset_discovers_only_valid_full_committed_calendar_months():
    from src.invoice_app.services.monthly_billing import (
        build_monthly_billing_dataset,
        month_label,
    )

    dataset = build_monthly_billing_dataset(_tabs())

    assert len(dataset.periods) == 1
    assert month_label(dataset.periods[0]) == "August 2026"


@pytest.mark.parametrize(
    ("commit_status", "validation_status"),
    (("READY", "PASSED"), ("COMMITTED", "FAILED")),
)
def test_uncommitted_or_failed_month_is_not_selectable(commit_status, validation_status):
    from src.invoice_app.services.monthly_billing import build_monthly_billing_dataset

    rows = (
        _monthly_row(commit_status=commit_status, validation_status=validation_status),
        _monthly_row(
            commit_status=commit_status,
            validation_status=validation_status,
            record_type="SKU",
        ),
    )

    assert build_monthly_billing_dataset(_tabs(rows)).periods == ()


def test_committed_non_calendar_month_fails_closed():
    from src.invoice_app.services.monthly_billing import (
        MonthlyBillingError,
        build_monthly_billing_dataset,
    )

    rows = (
        _monthly_row(period_to=date(2026, 8, 30)),
        _monthly_row(period_to=date(2026, 8, 30), record_type="SKU"),
    )

    with pytest.raises(MonthlyBillingError, match="full calendar month"):
        build_monthly_billing_dataset(_tabs(rows))


def test_duplicate_committed_monthly_batches_fail_closed():
    from src.invoice_app.services.monthly_billing import (
        MonthlyBillingError,
        build_monthly_billing_dataset,
    )

    rows = (
        _monthly_row(),
        _monthly_row(record_type="SKU"),
        _monthly_row(batch_id="monthly-august-revision", file_hash="r" * 64),
        _monthly_row(
            batch_id="monthly-august-revision",
            file_hash="r" * 64,
            record_type="SKU",
        ),
    )

    with pytest.raises(MonthlyBillingError, match="DUPLICATE MONTHLY STATEMENT MONTH"):
        build_monthly_billing_dataset(_tabs(rows))


def test_monthly_34_17_14_contract_and_absent_weekly_only_fields_are_accepted():
    from src.invoice_app.services.monthly_billing import build_monthly_billing_dataset

    assert len(MONTHLY_STATEMENT_DATA_HEADERS) == 34
    assert len(MONTHLY_STATEMENT_FINANCIAL_COMPONENT_HEADERS) == 17
    assert len(MONTHLY_STATEMENT_SUMMARY_HEADERS) == 14
    assert build_monthly_billing_dataset(_tabs()).periods


def test_monthly_report_reuses_existing_product_and_financial_calculations():
    from src.invoice_app.services.monthly_billing import (
        build_monthly_billing_dataset,
        build_monthly_billing_report,
    )
    from src.invoice_app.services.weekly_billing import build_weekly_billing_report

    dataset = build_monthly_billing_dataset(_tabs())
    report = build_monthly_billing_report(dataset, dataset.periods[0])

    assert report == build_weekly_billing_report(dataset, dataset.periods[0])
    assert report.product_summary.total_amount == Decimal("8.00")
    assert report.product_summary.product_rows[0].sku_code == "9555208000001"
    assert report.financial_summary.export_ready is True
    assert [control.passed for control in report.financial_summary.controls] == [
        True,
        True,
        True,
        True,
    ]


def test_monthly_missing_invoice_evidence_fails_closed_without_silent_exclusion():
    from src.invoice_app.services.monthly_billing import (
        MonthlyBillingError,
        build_monthly_billing_dataset,
        build_monthly_billing_report,
    )

    missing_order_id = "260823596770U3"
    rows = (
        _monthly_row(order_id=missing_order_id),
        _monthly_row(record_type="SKU"),
    )
    dataset = build_monthly_billing_dataset(_tabs(rows, orders=(), items=()))

    with pytest.raises(MonthlyBillingError, match=missing_order_id):
        build_monthly_billing_report(dataset, dataset.periods[0])


def test_monthly_workbook_reuses_final_product_summary_and_has_three_sheets():
    from src.invoice_app.services.monthly_billing import (
        build_monthly_billing_dataset,
        build_monthly_billing_report,
    )
    from src.invoice_app.services.weekly_billing_export import export_weekly_billing_report

    dataset = build_monthly_billing_dataset(_tabs())
    report = build_monthly_billing_report(dataset, dataset.periods[0])
    workbook = openpyxl.load_workbook(
        BytesIO(export_weekly_billing_report(report, product_master_records=()))
    )

    assert workbook.sheetnames == ["Product Summary", "Staging Data", "Financial Summary"]
    assert workbook["Product Summary"].cell(2, 2).value == "9555208000001"


def test_monthly_barcode_pdf_uses_monthly_wording_and_preserves_invalid_sku():
    from src.invoice_app.services.monthly_billing import (
        build_monthly_billing_dataset,
        build_monthly_billing_report,
    )
    from src.invoice_app.services.monthly_billing_export import (
        export_monthly_product_summary_barcode_table_pdf,
    )

    dataset = build_monthly_billing_dataset(_tabs())
    report = build_monthly_billing_report(dataset, dataset.periods[0])
    invalid_row = replace(report.product_summary.product_rows[0], sku_code="NOT-A-BARCODE")
    summary = replace(report.product_summary, product_rows=(invalid_row,))
    pdf = export_monthly_product_summary_barcode_table_pdf(summary)
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf)).pages)

    assert "MONTHLY BILLING PRODUCT SUMMARY" in text
    assert "Source: Shopee Monthly Statement" in text
    assert "NOT-A-BARCODE" in text
    assert "Barcode unavailable" in text


def test_sg_cannot_construct_the_monthly_my_reader():
    from src.invoice_app.services.market_context import MarketConfigurationUnavailable

    with pytest.raises(MarketConfigurationUnavailable):
        UAT2DataSettings(market_context=SHOPEE_SG).create_monthly_billing_reader()
