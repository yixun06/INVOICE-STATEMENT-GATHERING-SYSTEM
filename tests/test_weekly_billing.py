from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest

from src.invoice_app.domain.historical_invoice import (
    CanonicalInvoiceItem,
    CanonicalInvoiceOrder,
)
from src.invoice_app.domain.weekly_billing import (
    ActualSellingAmountBasis,
    BillingPeriod,
)
from src.invoice_app.repositories.google_sheets_historical_invoice_repository import (
    _serialize_item,
    _serialize_order,
)
from src.invoice_app.services.uat2_persistence_schema import (
    INVOICE_ITEMS_HEADERS,
    INVOICE_ITEMS_TAB,
    INVOICE_ORDERS_HEADERS,
    INVOICE_ORDERS_TAB,
    STATEMENT_DATA_HEADERS,
    STATEMENT_DATA_TAB,
    STATEMENT_FINANCIAL_COMPONENT_HEADERS,
    STATEMENT_FINANCIAL_COMPONENTS_TAB,
    STATEMENT_SUMMARY_HEADERS,
    STATEMENT_SUMMARY_TAB,
)
from src.invoice_app.services.weekly_billing import (
    WeeklyBillingDataset,
    WeeklyBillingError,
    build_weekly_billing_dataset,
    build_weekly_billing_summary,
)
from src.invoice_app.services.weekly_billing_export import (
    PRODUCT_SUMMARY_HEADERS,
    export_weekly_billing_summary,
)


PERIOD = BillingPeriod(
    date(2026, 8, 31),
    date(2026, 9, 6),
    "batch-golden",
    "a" * 64,
)


def _order(order_id: str) -> CanonicalInvoiceOrder:
    return CanonicalInvoiceOrder(
        platform="Shopee",
        order_id=order_id,
        first_imported_at=datetime(2026, 9, 7, tzinfo=timezone.utc),
    )


def _item(
    order_id: str,
    item_index: int,
    *,
    nav: str = "5000001",
    sku: str = "SKU-1",
    name: str = "Product One",
    variation: str | None = None,
    price: str = "10.00",
    quantity: int = 1,
    subtotal: str | None = "8.00",
    promotion_group: str | None = None,
    promotion_total: str | None = None,
) -> CanonicalInvoiceItem:
    return CanonicalInvoiceItem(
        platform="Shopee",
        order_id=order_id,
        item_index=item_index,
        nav=nav,
        seller_sku=sku,
        product_name=name,
        variation=variation,
        unit_price=Decimal(price),
        quantity=quantity,
        line_subtotal=Decimal(subtotal) if subtotal is not None else None,
        promotion_group_id=promotion_group,
        promotion_label=("Any promotion" if promotion_group else None),
        source_group_total=(
            Decimal(promotion_total) if promotion_total is not None else None
        ),
    )


def _dataset(
    *items: CanonicalInvoiceItem,
    order_ids: tuple[str, ...] | None = None,
) -> WeeklyBillingDataset:
    selected = order_ids or tuple(dict.fromkeys(item.order_id for item in items))
    return WeeklyBillingDataset(
        periods=(PERIOD,),
        order_ids_by_batch={PERIOD.statement_batch_id: selected},
        orders={order_id: _order(order_id) for order_id in selected},
        items=tuple(items),
    )


def _statement_row(
    *,
    batch_id: str = "batch-golden",
    file_hash: str = "a" * 64,
    period_from: str = "2026-08-31",
    period_to: str = "2026-09-06",
    commit_status: str = "COMMITTED",
    record_type: str = "ORDER",
    sequence_no: str = "1",
    order_id: str = "ORDER-1",
    order_count: str = "1",
) -> tuple[str, ...]:
    values = {header: "" for header in STATEMENT_DATA_HEADERS}
    values.update(
        {
            "statement_batch_id": batch_id,
            "record_type": record_type,
            "sequence_no": sequence_no,
            "platform": "Shopee",
            "statement_file_hash": file_hash,
            "statement_period_from": period_from,
            "statement_period_to": period_to,
            "statement_order_count": order_count,
            "validation_status": "PASSED",
            "commit_status": commit_status,
            "order_id": order_id if record_type == "ORDER" else "",
        }
    )
    return tuple(values[header] for header in STATEMENT_DATA_HEADERS)


def _tabs(
    statement_rows: tuple[tuple[str, ...], ...],
    *,
    orders: tuple[CanonicalInvoiceOrder, ...] = (),
    items: tuple[CanonicalInvoiceItem, ...] = (),
):
    return {
        STATEMENT_DATA_TAB: (STATEMENT_DATA_HEADERS, *statement_rows),
        INVOICE_ORDERS_TAB: (
            INVOICE_ORDERS_HEADERS,
            *(_serialize_order(order) for order in orders),
        ),
        INVOICE_ITEMS_TAB: (
            INVOICE_ITEMS_HEADERS,
            *(_serialize_item(item) for item in items),
        ),
        STATEMENT_FINANCIAL_COMPONENTS_TAB: (
            STATEMENT_FINANCIAL_COMPONENT_HEADERS,
        ),
        STATEMENT_SUMMARY_TAB: (STATEMENT_SUMMARY_HEADERS,),
    }


def test_period_discovery_includes_only_committed_and_newest_first():
    old = _statement_row(
        batch_id="old",
        file_hash="b" * 64,
        period_from="2026-08-17",
        period_to="2026-08-23",
    )
    uncommitted = _statement_row(
        batch_id="draft",
        file_hash="c" * 64,
        period_from="2026-09-07",
        period_to="2026-09-13",
        commit_status="READY",
    )
    current = _statement_row()
    dataset = build_weekly_billing_dataset(
        _tabs((old, uncommitted, current), orders=(_order("ORDER-1"),))
    )

    assert [period.statement_batch_id for period in dataset.periods] == [
        "batch-golden",
        "old",
    ]
    assert "draft" not in dataset.order_ids_by_batch


def test_same_period_multiple_committed_batches_fails_closed():
    duplicate = _statement_row(batch_id="another", file_hash="b" * 64)
    with pytest.raises(
        WeeklyBillingError,
        match="BUSINESS DECISION REQUIRED — DUPLICATE STATEMENT PERIOD",
    ):
        build_weekly_billing_dataset(_tabs((_statement_row(), duplicate)))


def test_arbitrary_period_is_not_supported():
    arbitrary = replace(PERIOD, statement_batch_id="not-committed")
    with pytest.raises(WeeklyBillingError, match="not an existing COMMITTED"):
        build_weekly_billing_summary(_dataset(_item("ORDER-1", 0)), arbitrary)


def test_selected_statement_orders_define_complete_source_population():
    selected = _item("ORDER-1", 0, quantity=2, subtotal="16.00")
    unrelated = _item("ORDER-2", 0, subtotal="99.00")
    dataset = _dataset(selected, unrelated, order_ids=("ORDER-1",))

    summary = build_weekly_billing_summary(dataset, PERIOD)

    assert summary.order_count == 1
    assert summary.invoice_item_count == 1
    assert {(item.order_id, item.item_index) for item in summary.source_items} == {
        ("ORDER-1", 0)
    }
    assert summary.total_quantity == 2
    assert summary.total_amount == Decimal("16.00")


def test_normal_items_with_different_seller_skus_remain_separate():
    summary = build_weekly_billing_summary(
        _dataset(
            _item("ORDER-1", 0, sku="SKU-A", quantity=2, subtotal="16.00"),
            _item("ORDER-2", 0, sku="SKU-B", quantity=3, subtotal="21.00"),
        ),
        PERIOD,
    )

    assert len(summary.product_rows) == 2
    assert sum(row.quantity for row in summary.product_rows) == 5
    assert sum(row.amount for row in summary.product_rows) == Decimal("37.00")
    assert sum(row.discount_amount for row in summary.product_rows) == Decimal("13.00")
    assert all(row.uom is None for row in summary.product_rows)
    assert all(row.discount_percent is None for row in summary.product_rows)
    assert summary.source_items[0].actual_selling_amount_basis is ActualSellingAmountBasis.DIRECT


@pytest.mark.parametrize(
    ("case", "first", "second", "expected_rows", "description"),
    (
        (
            "L-01 keeps conflicting tea titles separate",
            _item("ORDER-1", 0, nav="5004322", sku="9555208108580", name="Rose Tea", variation=None, price="18.90"),
            _item("ORDER-2", 0, nav="5004322", sku="9555208108580", name="Jasmine Tea Rose Tea", variation=None, price="18.90"),
            2,
            None,
        ),
        (
            "L-02 merges approved honey variations",
            _item("ORDER-1", 0, nav="3000209", sku="9555208107347", name="Pre-Order Simply Natural Fresh Raw Honey Malaysia [Madu Asli Segar]", variation="1KG", price="54.90"),
            _item("ORDER-2", 0, nav="3000209", sku="9555208107347", name="Simply Natural Fresh Raw Honey Malaysia [Madu Asli Segar]", variation="Fresh Raw Honey 1kg", price="54.90"),
            1,
            "Simply Natural Fresh Raw Honey Malaysia [Madu Asli Segar] | 1KG",
        ),
        (
            "L-03 keeps Coconut Sugar and Normal separate",
            _item("ORDER-1", 0, nav="5003317", sku="9555208104926", variation="Coconut Sugar", price="13.90"),
            _item("ORDER-2", 0, nav="5003317", sku="9555208104926", variation="Normal", price="13.90"),
            2,
            None,
        ),
        (
            "L-04 keeps Baby Noodles ambiguity separate",
            _item("ORDER-1", 0, nav="4007457", sku="9555208017721", name="Baby Thin Noodle", variation="Rainbow Baby Noodles", price="12.90"),
            _item("ORDER-2", 0, nav="4007457", sku="9555208017721", name="Rainbow Noodles", variation=None, price="12.90"),
            2,
            None,
        ),
        (
            "L-05 keeps Good Fibre variants separate",
            _item("ORDER-1", 0, nav="5003397", sku="9555208105169", name="Good Fibre", variation="1 Box", price="79.90"),
            _item("ORDER-2", 0, nav="5003397", sku="9555208105169", name="Good Fibre Plus+ Zero Sugar", variation="1box (Normal sugar)", price="79.90"),
            2,
            None,
        ),
        (
            "L-06 merges approved Sweet Potato Mee Sua variations",
            _item("ORDER-1", 0, nav="5000165", sku="9555208103158", name="Simply Natural Organic Mee Sua 200g Malaysia", variation=None, price="8.50"),
            _item("ORDER-2", 0, nav="5000165", sku="9555208103158", name="Simply Natural Organic Handmade Sweet Potato Mee Sua 200g Malaysia", variation="Sweet Potato Mee Sua", price="8.50"),
            1,
            "Simply Natural Organic Handmade Sweet Potato Mee Sua 200g Malaysia | Sweet Potato Mee Sua",
        ),
        (
            "L-07 keeps historical price changes separate",
            _item("ORDER-1", 0, nav="060328", sku="9555208105145-1Lter", variation="1000ml", price="42.90"),
            _item("ORDER-2", 0, nav="060328", sku="9555208105145-1Lter", variation="1000ml", price="43.90"),
            2,
            None,
        ),
        (
            "L-08 keeps Buy 5 free 1 and Bundle Set separate",
            _item("ORDER-1", 0, nav="3000573", sku="9555208106944-6", variation="Buy 5 free 1", price="36.00"),
            _item("ORDER-2", 0, nav="3000573", sku="9555208106944-6", variation="Bundle Set", price="36.00"),
            2,
            None,
        ),
    ),
    ids=lambda value: value if isinstance(value, str) else None,
)
def test_product_owner_approved_listing_identity_cases(
    case, first, second, expected_rows, description,
):
    summary = build_weekly_billing_summary(_dataset(first, second), PERIOD)

    assert len(summary.product_rows) == expected_rows
    assert sum(row.quantity for row in summary.product_rows) == 2
    assert sum(row.amount for row in summary.product_rows) == Decimal("16.00")
    if description is not None:
        assert summary.product_rows[0].product_name == description


@pytest.mark.parametrize(
    ("first_name", "second_name", "expected_description"),
    (
        ("Gluten F ree", "Gluten Free", "Gluten Free"),
        ("Plant-Ba sed Oil", "Plant-Based Oil", "Plant-Based Oil"),
        ("Best Snacks", "Be st Snacks", "Best Snacks"),
        ("Pre-Order Product One", "Product One", "Product One"),
        ("中文 名称", "中文名称", "中文名称"),
    ),
)
def test_approved_technical_title_noise_merges(
    first_name, second_name, expected_description,
):
    summary = build_weekly_billing_summary(
        _dataset(
            _item("ORDER-1", 0, name=first_name, variation="Original"),
            _item("ORDER-2", 0, name=second_name, variation="Original"),
        ),
        PERIOD,
    )

    assert len(summary.product_rows) == 1
    assert summary.product_rows[0].product_name == f"{expected_description} | Original"


def test_same_identity_with_different_historical_price_remains_two_rows():
    first = replace(
        _item("ORDER-1", 0, price="10.00", subtotal="8.00"),
        seller_sku=None,
        sku_missing_in_source=True,
    )
    second = replace(
        _item("ORDER-2", 0, price="12.00", subtotal="9.00"),
        seller_sku=None,
        sku_missing_in_source=True,
    )
    summary = build_weekly_billing_summary(
        _dataset(first, second),
        PERIOD,
    )

    assert [row.unit_price for row in summary.product_rows] == [
        Decimal("10.00"),
        Decimal("12.00"),
    ]


@pytest.mark.parametrize(
    ("field", "value", "message"),
    (
        ("nav", None, "nav"),
        ("product_name", None, "product_name"),
        ("unit_price", None, "historical PM Unit Price"),
        ("quantity", None, "quantity"),
        ("quantity", 0, "quantity"),
    ),
)
def test_missing_required_billing_fact_fails_loudly(field, value, message):
    with pytest.raises(WeeklyBillingError, match=message):
        build_weekly_billing_summary(
            _dataset(replace(_item("ORDER-1", 0), **{field: value})), PERIOD
        )


def test_missing_normal_line_subtotal_fails_loudly():
    with pytest.raises(WeeklyBillingError, match="line_subtotal"):
        build_weekly_billing_summary(
            _dataset(_item("ORDER-1", 0, subtotal=None)), PERIOD
        )


def test_source_missing_seller_sku_is_preserved_and_does_not_block_billing():
    item = replace(
        _item(
            "260901TSAE91UF",
            0,
            nav="5004722",
            name="Simply Natural 3 Treasure Powder 200g",
            price="20.90",
            quantity=2,
            subtotal="41.80",
        ),
        seller_sku=None,
        sku_missing_in_source=True,
    )

    peer = _item(
        "ORDER-2",
        0,
        nav="5004722",
        sku="SOURCE-SKU-EVIDENCE",
        name="Simply Natural 3 Treasure Powder 200g",
        price="20.90",
        quantity=1,
        subtotal="20.90",
    )
    summary = build_weekly_billing_summary(_dataset(item, peer), PERIOD)

    assert any(source.seller_sku is None for source in summary.source_items)
    assert len(summary.product_rows) == 2
    assert {row.nav for row in summary.product_rows} == {"5004722"}
    assert sum(row.quantity for row in summary.product_rows) == 3
    assert sum(row.amount for row in summary.product_rows) == Decimal("62.70")
    assert sum(row.discount_amount for row in summary.product_rows) == Decimal("0.00")


def test_negative_discount_is_preserved_not_clamped():
    summary = build_weekly_billing_summary(
        _dataset(
            _item(
                "2609045MDX0J9G",
                0,
                nav="3000573",
                name="Simply Natural Baked Brown Rice Cracker 140g",
                price="36.00",
                quantity=1,
                subtotal="39.90",
            )
        ),
        PERIOD,
    )

    assert summary.product_rows[0].discount_amount == Decimal("-3.90")
    assert summary.total_discount_amount == Decimal("-3.90")
    assert summary.product_rows[0].unit_price - Decimal("-3.90") == Decimal(
        "39.90"
    )


def test_same_price_promotion_uses_source_total_and_exact_largest_remainder():
    members = (
        _item(
            "ORDER-1", 0, sku="SKU-A", nav="NAV-A", price="20.00",
            quantity=1, subtotal="20.00", promotion_group="P1",
            promotion_total="45.80",
        ),
        _item(
            "ORDER-1", 1, sku="SKU-B", nav="NAV-B", price="20.00",
            quantity=2, subtotal="40.00", promotion_group="P1",
            promotion_total="45.80",
        ),
    )
    summary = build_weekly_billing_summary(_dataset(*members), PERIOD)

    assert [item.actual_selling_amount for item in summary.source_items] == [
        Decimal("15.27"),
        Decimal("30.53"),
    ]
    assert summary.promotion_amount_total == Decimal("45.80")
    assert summary.same_price_promotion_count == 1
    assert summary.mixed_price_promotion_count == 0


def test_golden_mixed_price_promotion_allocates_9753_and_9047_without_mutation():
    members = (
        _item(
            "260828FVPP5BTR", 0, sku="9555208107590", nav="5003908",
            price="138.00", subtotal="138.00",
            promotion_group="shopee-promotion:p1:section1:group1",
            promotion_total="188.00",
        ),
        _item(
            "260828FVPP5BTR", 1, sku="9555208105220", nav="5003322",
            price="128.00", subtotal="128.00",
            promotion_group="shopee-promotion:p1:section1:group1",
            promotion_total="188.00",
        ),
    )
    original = deepcopy(members)

    summary = build_weekly_billing_summary(_dataset(*members), PERIOD)
    amounts = {item.seller_sku: item.actual_selling_amount for item in summary.source_items}

    assert amounts == {
        "9555208107590": Decimal("97.53"),
        "9555208105220": Decimal("90.47"),
    }
    assert sum(amounts.values()) == Decimal("188.00")
    discounts = {
        item.seller_sku: item.discount_amount for item in summary.source_items
    }
    assert discounts == {
        "9555208107590": Decimal("40.47"),
        "9555208105220": Decimal("37.53"),
    }
    assert summary.total_standard_amount == Decimal("266.00")
    assert summary.total_discount_amount == Decimal("78.00")
    assert summary.total_amount == Decimal("188.00")
    assert all(
        row.unit_price * row.quantity - row.discount_amount == row.amount
        for row in summary.product_rows
    )
    assert summary.mixed_price_promotion_count == 1
    assert all(
        item.actual_selling_amount_basis
        is ActualSellingAmountBasis.PROMOTION_MIXED_PRICE_WEIGHTED_ALLOCATED
        for item in summary.source_items
    )
    assert members == original


def test_largest_remainder_tie_break_is_stable_under_input_permutation():
    members = (
        _item(
            "ORDER-1", 0, sku="SKU-B", nav="NAV-B", price="10.00",
            subtotal="10.00", promotion_group="P1", promotion_total="10.01",
        ),
        _item(
            "ORDER-1", 1, sku="SKU-A", nav="NAV-A", price="10.00",
            subtotal="10.00", promotion_group="P1", promotion_total="10.01",
        ),
    )
    forward = build_weekly_billing_summary(_dataset(*members), PERIOD)
    reverse = build_weekly_billing_summary(_dataset(*reversed(members)), PERIOD)

    def allocations(summary):
        return {item.seller_sku: item.actual_selling_amount for item in summary.source_items}

    assert allocations(forward) == allocations(reverse) == {
        "SKU-A": Decimal("5.01"),
        "SKU-B": Decimal("5.00"),
    }


def test_promotion_group_missing_or_conflicting_total_fails_closed():
    missing = _item(
        "ORDER-1", 0, promotion_group="P1", promotion_total=None
    )
    with pytest.raises(WeeklyBillingError, match="missing or conflicting"):
        build_weekly_billing_summary(_dataset(missing), PERIOD)

    conflicting = (
        _item("ORDER-1", 0, promotion_group="P1", promotion_total="10.00"),
        _item("ORDER-1", 1, promotion_group="P1", promotion_total="11.00"),
    )
    with pytest.raises(WeeklyBillingError, match="missing or conflicting"):
        build_weekly_billing_summary(_dataset(*conflicting), PERIOD)


def test_v2_group_order_is_included_without_statement_allocation_or_refund_netting():
    statement = (
        _statement_row(order_id="ORDER-1"),
        _statement_row(
            record_type="SKU", sequence_no="2", order_id="", order_count="1"
        ),
        _statement_row(
            record_type="ADJUSTMENT", sequence_no="3", order_id="", order_count="1"
        ),
    )
    item = replace(
        _item("ORDER-1", 0, subtotal="25.00"),
        statement_refund_amount=Decimal("-10.00"),
        statement_net_selling_amount=None,
    )
    dataset = build_weekly_billing_dataset(
        _tabs(statement, orders=(_order("ORDER-1"),), items=(item,))
    )
    summary = build_weekly_billing_summary(dataset, dataset.periods[0])

    assert summary.invoice_item_count == 1
    assert summary.total_amount == Decimal("25.00")
    assert summary.source_items[0].item_index == 0


def test_duplicate_item_and_missing_order_item_are_blocked():
    duplicate = _item("ORDER-1", 0)
    with pytest.raises(WeeklyBillingError, match="counted twice"):
        build_weekly_billing_summary(_dataset(duplicate, duplicate), PERIOD)
    with pytest.raises(WeeklyBillingError, match="have no Invoice_Items"):
        build_weekly_billing_summary(
            _dataset(_item("ORDER-1", 0), order_ids=("ORDER-1", "ORDER-2")),
            PERIOD,
        )


def test_quantity_control_rejects_a_corrupt_aggregation(monkeypatch):
    from src.invoice_app.services import weekly_billing as module

    original = module._aggregate_source_items

    def corrupt(source_items):
        rows = original(source_items)
        return (replace(rows[0], quantity=rows[0].quantity + 1), *rows[1:])

    monkeypatch.setattr(module, "_aggregate_source_items", corrupt)
    with pytest.raises(WeeklyBillingError, match="Quantity control"):
        build_weekly_billing_summary(_dataset(_item("ORDER-1", 0)), PERIOD)


def test_selling_amount_control_rejects_a_corrupt_aggregation(monkeypatch):
    from src.invoice_app.services import weekly_billing as module

    original = module._aggregate_source_items

    def corrupt(source_items):
        rows = original(source_items)
        return (
            replace(
                rows[0],
                amount=rows[0].amount + Decimal("0.01"),
            ),
            *rows[1:],
        )

    monkeypatch.setattr(module, "_aggregate_source_items", corrupt)
    with pytest.raises(WeeklyBillingError, match="Amount control"):
        build_weekly_billing_summary(_dataset(_item("ORDER-1", 0)), PERIOD)


def test_discount_control_rejects_a_corrupt_aggregation(monkeypatch):
    from src.invoice_app.services import weekly_billing as module

    original = module._aggregate_source_items

    def corrupt(source_items):
        rows = original(source_items)
        return (
            replace(
                rows[0],
                discount_amount=rows[0].discount_amount + Decimal("0.01"),
            ),
            *rows[1:],
        )

    monkeypatch.setattr(module, "_aggregate_source_items", corrupt)
    with pytest.raises(WeeklyBillingError, match="Disc Amt control"):
        build_weekly_billing_summary(_dataset(_item("ORDER-1", 0)), PERIOD)


def test_excel_uses_exact_summary_rows_numeric_money_and_deterministic_order():
    summary = build_weekly_billing_summary(
        _dataset(
            _item("ORDER-2", 0, nav="NAV-B", sku="SKU-B", name="Beta"),
            _item("ORDER-1", 0, nav="NAV-A", sku="SKU-A", name="Alpha"),
        ),
        PERIOD,
    )
    workbook = openpyxl.load_workbook(
        BytesIO(export_weekly_billing_summary(summary)), data_only=True
    )
    sheet = workbook["Product Summary"]
    values = list(sheet.iter_rows(values_only=True))

    assert values[0] == PRODUCT_SUMMARY_HEADERS
    assert [row[:4] for row in values[1:]] == [
        (1, "NAV-A", "Alpha", 1),
        (2, "NAV-B", "Beta", 1),
    ]
    assert all(sheet.cell(row=row, column=6).data_type == "n" for row in (2, 3))
    assert all(sheet.cell(row=row, column=8).data_type == "n" for row in (2, 3))
    assert all(sheet.cell(row=row, column=9).data_type == "n" for row in (2, 3))
    assert all(row[4] is None and row[6] is None for row in values[1:])
    assert values[1:] == [
        (1, "NAV-A", "Alpha", 1, None, 10, None, 2, 8),
        (2, "NAV-B", "Beta", 1, None, 10, None, 2, 8),
    ]
    assert sum(row[3] for row in values[1:]) == summary.total_quantity
    assert Decimal(str(sum(row[7] for row in values[1:]))) == summary.total_discount_amount
    assert Decimal(str(sum(row[8] for row in values[1:]))) == summary.total_amount
