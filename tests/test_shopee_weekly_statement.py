from __future__ import annotations

from collections import Counter
from dataclasses import replace
from datetime import date
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from src.invoice_app.parsers.shopee_weekly_statement_parser import (
    INCOME_COMPONENT_COLUMNS,
    INCOME_REQUIRED_COLUMNS,
    _parse_income_rows,
    parse_shopee_weekly_statement,
    WeeklyStatementParseError,
    INVALID_REQUIRED_HEADER,
    MISSING_REQUIRED_WORKSHEET,
)
from src.invoice_app.services.shopee_weekly_statement_service import (
    ALREADY_IMPORTED,
    NEEDS_REVIEW,
    READY_TO_COMMIT,
    REJECTED,
    StatementReference,
    stage_parsed_shopee_weekly_statement,
    stage_shopee_weekly_statement,
    validate_shopee_weekly_statement,
)


ROOT = Path(__file__).resolve().parents[1]
SAMPLE = (
    ROOT / "tests" / "fixtures" / "shopee_weekly_statement"
    / "Income.released.my.20260427_20260503.xlsx"
)


class FakeUploadedFile:
    def __init__(self, name: str, data: bytes):
        self.name = name
        self._data = data

    def getvalue(self) -> bytes:
        return self._data


@pytest.fixture(scope="module")
def parsed_sample():
    before = SAMPLE.read_bytes()
    statement = parse_shopee_weekly_statement(
        FakeUploadedFile(SAMPLE.name, before)
    )
    after = SAMPLE.read_bytes()
    assert before == after
    return statement


def test_native_shopee_export_dimension_fallback_and_contract(parsed_sample):
    with ZipFile(SAMPLE) as archive:
        with archive.open("xl/worksheets/sheet5.xml") as income_xml:
            xml_prefix = income_xml.read(1024).decode("utf-8")

    assert '<dimension ref="A1"' in xml_prefix
    assert parsed_sample.dimension_fallback_sheets == ("Income",)
    assert parsed_sample.statement_period_from == date(2026, 4, 27)
    assert parsed_sample.statement_period_to == date(2026, 5, 3)
    assert len(parsed_sample.order_rows) == 505
    assert len(parsed_sample.sku_rows) == 820
    assert len(parsed_sample.service_fee_details) == 501
    assert len(parsed_sample.shipping_fee_discrepancies) == 14
    assert len(parsed_sample.adjustments) == 3
    assert parsed_sample.summary_total_released == Decimal("20599.26")
    assert parsed_sample.adjustment_control_total == Decimal("-126.63")
    assert parsed_sample.source_value_issues == ()
    assert parsed_sample.file_hash == sha256(SAMPLE.read_bytes()).hexdigest()


def test_income_rows_before_late_native_header_are_preserved_with_source_numbers():
    header = tuple(INCOME_REQUIRED_COLUMNS)
    values = {name: Decimal("0.00") for name in INCOME_COMPONENT_COLUMNS}
    values.update(
        {
            "Sequence No.": "1",
            "View By": "Order",
            "Order ID": "ORDER-BEFORE-HEADER",
            "Product ID": "",
            "Product Name": "",
            "Order Creation Date": date(2026, 8, 1),
            "Payout Completed Date": date(2026, 8, 31),
            "Release Channel": "Seller Wallet",
            "Order Type": "Normal Order",
            "Total Released Amount (RM)": Decimal("10.00"),
            "Product Price": Decimal("10.00"),
        }
    )
    source_row = tuple(values.get(name) for name in header)
    group_heading = [None] * len(header)
    group_heading[0] = "Order Info"
    group_heading[11] = "Released Amount Details"
    issues = []

    parsed = _parse_income_rows(
        [
            tuple(group_heading),
            source_row,
            header,
            replace_tuple(source_row, 0, "2"),
        ],
        issues,
    )

    assert [row.source_row_number for row in parsed] == [2, 4]
    assert [row.order_id for row in parsed] == [
        "ORDER-BEFORE-HEADER",
        "ORDER-BEFORE-HEADER",
    ]
    assert issues == []


def test_malformed_income_row_before_late_header_is_not_silently_discarded():
    header = tuple(INCOME_REQUIRED_COLUMNS)
    malformed = [None] * len(header)
    malformed[header.index("Product Name")] = "Unexpected source row"
    issues = []

    parsed = _parse_income_rows([tuple(malformed), header], issues)

    assert len(parsed) == 1
    assert parsed[0].source_row_number == 1
    assert [issue.code for issue in issues] == ["unrecognized_income_row"]


def replace_tuple(values, index, value):
    result = list(values)
    result[index] = value
    return tuple(result)


def test_native_summary_preserves_approved_lines_hierarchy_currency_and_money(parsed_sample):
    lines = {line.native_label: line for line in parsed_sample.summary_lines}

    assert len(parsed_sample.summary_lines) == 32
    assert "Income Summary" not in lines
    assert lines["1. Total Revenue"].line_type == "TOTAL"
    assert lines["Merchandise Subtotal"].line_type == "SUBTOTAL"
    assert lines["Original product price"].line_type == "DETAIL"
    assert lines["Other Reference Values"].line_type == "SECTION_HEADER"
    assert lines["Shipping Fee Promotion by Seller"].line_type == "REFERENCE"
    assert lines["Original product price"].parent_source_row_number == lines["Merchandise Subtotal"].statement_source_row_number
    assert lines["Shipping Fee Promotion by Seller"].parent_source_row_number == lines["Other Reference Values"].statement_source_row_number
    assert lines["Other Reference Values"].component_amount is None
    assert lines["Rebate Provided by Shopee"].component_amount == Decimal("0.00")
    assert lines["Your Seller product promotion"].component_amount < 0
    assert all(line.currency == "RM" for line in parsed_sample.summary_lines)


def _summary_mutation(mutator):
    workbook = load_workbook(BytesIO(SAMPLE.read_bytes()))
    mutator(workbook["Summary"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _workbook_mutation(mutator) -> bytes:
    workbook = load_workbook(BytesIO(SAMPLE.read_bytes()))
    mutator(workbook)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_valid_statement_without_adjustment_sheet_is_accepted_as_no_events():
    source = _workbook_mutation(
        lambda workbook: workbook.remove(workbook["Adjustment"])
    )

    parsed = parse_shopee_weekly_statement(
        source,
        source_filename="statement-without-adjustment.xlsx",
    )

    assert parsed.adjustments == ()
    assert parsed.adjustment_control_total == Decimal("0.00")
    assert parsed.adjustment_footer_total is None
    assert "Adjustment" not in parsed.dimension_fallback_sheets
    assert validate_shopee_weekly_statement(parsed) == ()


def test_present_but_malformed_adjustment_sheet_still_fails_closed():
    source = _workbook_mutation(
        lambda workbook: setattr(
            workbook["Adjustment"]["A13"],
            "value",
            "Unsupported Sequence Header",
        )
    )

    with pytest.raises(WeeklyStatementParseError) as error:
        parse_shopee_weekly_statement(
            source,
            source_filename="malformed-adjustment.xlsx",
        )

    assert error.value.code == INVALID_REQUIRED_HEADER
    assert error.value.sheet_name == "Adjustment"


def test_present_empty_valid_adjustment_sheet_keeps_existing_contract():
    def empty_adjustment(workbook):
        sheet = workbook["Adjustment"]
        sheet.delete_rows(14, 3)
        sheet["E8"] = Decimal("0.00")
        sheet["E15"] = Decimal("0.00")

    parsed = parse_shopee_weekly_statement(
        _workbook_mutation(empty_adjustment),
        source_filename="empty-adjustment.xlsx",
    )

    assert parsed.adjustments == ()
    assert parsed.adjustment_control_total == Decimal("0.00")
    assert parsed.adjustment_footer_total == Decimal("0.00")
    assert validate_shopee_weekly_statement(parsed) == ()


def test_missing_non_adjustment_sheet_remains_structured_source_failure():
    source = _workbook_mutation(
        lambda workbook: workbook.remove(workbook["Income"])
    )

    staged = stage_shopee_weekly_statement(
        source,
        source_filename="statement-without-income.xlsx",
    )

    assert staged.result == REJECTED
    assert staged.statement is None
    assert staged.source_error is not None
    assert staged.source_error.code == MISSING_REQUIRED_WORKSHEET
    assert staged.source_error.sheet_name == "Income"
    assert "Adjustment" in staged.source_error.available_sheets


@pytest.mark.parametrize(
    "mutator, message",
    [
        (lambda sheet: setattr(sheet["C16"], "value", None), "must contain exactly one amount"),
        (lambda sheet: setattr(sheet["B16"], "value", "Unknown Seller Charge"), "Unsupported meaningful Summary line"),
        (lambda sheet: setattr(sheet["B16"], "value", None), "Unsupported meaningful Summary line"),
        (lambda sheet: (setattr(sheet["A16"], "value", sheet["B16"].value), setattr(sheet["B16"], "value", None)), "hierarchy is unsupported"),
    ],
)
def test_native_summary_unsupported_or_missing_source_fails_closed(mutator, message):
    with pytest.raises(WeeklyStatementParseError, match=message):
        parse_shopee_weekly_statement(
            _summary_mutation(mutator), source_filename=SAMPLE.name
        )


def test_native_summary_allows_row_shift_and_blank_spacers_without_fixed_coordinates():
    shifted = _summary_mutation(lambda sheet: sheet.insert_rows(14, 2))
    parsed = parse_shopee_weekly_statement(shifted, source_filename=SAMPLE.name)

    assert len(parsed.summary_lines) == 32
    assert next(
        line for line in parsed.summary_lines
        if line.native_label == "1. Total Revenue"
    ).statement_source_row_number == 16


def test_native_sample_requires_review_when_target_orders_are_unmatched(
    parsed_sample,
):
    staged = stage_parsed_shopee_weekly_statement(parsed_sample)

    assert staged.result == NEEDS_REVIEW
    assert staged.validation_issues == ()
    assert staged.eligible_for_future_atomic_commit is False
    assert staged.whole_statement_atomic is True
    assert staged.reconciliation_counts == {
        "Matched": 0,
        "Different": 0,
        "Estimated Only": 0,
        "Unmatched Order": 505,
        "Missing Comparison Evidence": 0,
    }
    assert Counter(item.status for item in staged.adjustment_reconciliations) == {
        "Unmatched Adjustment": 3
    }

def _replace_income_row(statement, original, replacement):
    rows = list(statement.income_rows)
    rows[rows.index(original)] = replacement
    return replace(statement, income_rows=tuple(rows))


def _validation_codes(statement):
    return {issue.code for issue in validate_shopee_weekly_statement(statement)}


def _august_statement_with_in_period_income_and_adjustments(parsed_sample):
    august_start = date(2026, 8, 1)
    return replace(
        parsed_sample,
        statement_period_from=august_start,
        statement_period_to=date(2026, 8, 31),
        income_rows=tuple(
            replace(row, payout_completed_date=august_start)
            for row in parsed_sample.income_rows
        ),
        adjustments=tuple(
            replace(
                row,
                adjustment_complete_date=august_start,
                payout_completed_date=august_start,
            )
            for row in parsed_sample.adjustments
        ),
    )


def test_adjustment_period_uses_complete_date_not_historical_payout_date(parsed_sample):
    statement = _august_statement_with_in_period_income_and_adjustments(parsed_sample)
    replacement_amount = Decimal("-17.67")
    adjustment_total_delta = replacement_amount - statement.adjustments[0].adjustment_amount
    adjustment = replace(
        statement.adjustments[0],
        adjustment_complete_date=date(2026, 8, 1),
        payout_completed_date=date(2026, 7, 27),
        adjustment_amount=replacement_amount,
    )
    statement = replace(
        statement,
        adjustments=(adjustment, *statement.adjustments[1:]),
        adjustment_control_total=statement.adjustment_control_total + adjustment_total_delta,
        adjustment_footer_total=statement.adjustment_footer_total + adjustment_total_delta,
    )

    assert "adjustment_complete_date_outside_statement_period" not in _validation_codes(statement)
    assert stage_parsed_shopee_weekly_statement(statement).result == NEEDS_REVIEW
    assert statement.adjustments[0].payout_completed_date == date(2026, 7, 27)


def test_adjustment_complete_date_outside_statement_period_still_fails(parsed_sample):
    statement = _august_statement_with_in_period_income_and_adjustments(parsed_sample)
    adjustment = replace(
        statement.adjustments[0],
        adjustment_complete_date=date(2026, 7, 31),
        payout_completed_date=date(2026, 8, 1),
    )
    statement = replace(statement, adjustments=(adjustment, *statement.adjustments[1:]))

    assert "adjustment_complete_date_outside_statement_period" in _validation_codes(statement)


def test_income_payout_period_validation_is_unchanged(parsed_sample):
    statement = _august_statement_with_in_period_income_and_adjustments(parsed_sample)
    income = replace(statement.income_rows[0], payout_completed_date=date(2026, 7, 31))
    statement = replace(statement, income_rows=(income, *statement.income_rows[1:]))

    assert "payout_date_outside_statement_period" in _validation_codes(statement)


def test_footer_like_income_row_is_preserved_as_one_source_quality_warning():
    header = tuple(INCOME_REQUIRED_COLUMNS)
    footer = [None] * len(header)
    footer[header.index("Product Name")] = "Order Income"
    issues = []

    parsed = _parse_income_rows([header, tuple(footer)], issues)

    assert len(parsed) == 1
    assert parsed[0].view_by == ""
    assert parsed[0].order_id == ""
    assert [issue.code for issue in issues] == ["unrecognized_income_row"]


def test_each_confirmed_blocking_financial_validation_detects_mismatch(parsed_sample):
    assert "order_total_vs_summary_mismatch" in _validation_codes(
        replace(
            parsed_sample,
            summary_total_released=parsed_sample.summary_total_released
            + Decimal("0.03"),
        )
    )

    order_row = parsed_sample.order_rows[0]
    components = dict(order_row.financial_components)
    components[INCOME_COMPONENT_COLUMNS[0]] += Decimal("0.03")
    assert "order_component_mismatch" in _validation_codes(
        _replace_income_row(
            parsed_sample,
            order_row,
            replace(order_row, financial_components=components),
        )
    )

    sku_row = parsed_sample.sku_rows[0]
    assert "sku_total_vs_order_mismatch" in _validation_codes(
        _replace_income_row(
            parsed_sample,
            sku_row,
            replace(
                sku_row,
                total_released_amount=sku_row.total_released_amount
                + Decimal("0.03"),
            ),
        )
    )

    detail = parsed_sample.service_fee_details[0]
    fee_components = dict(detail.components)
    first_fee = next(iter(fee_components))
    fee_components[first_fee] += Decimal("0.03")
    service_details = list(parsed_sample.service_fee_details)
    service_details[0] = replace(detail, components=fee_components)
    assert "service_fee_detail_mismatch" in _validation_codes(
        replace(parsed_sample, service_fee_details=tuple(service_details))
    )

    assert "adjustment_detail_vs_control_mismatch" in _validation_codes(
        replace(
            parsed_sample,
            adjustment_control_total=parsed_sample.adjustment_control_total
            + Decimal("0.03"),
        )
    )


def test_invalid_statement_period_is_blocking(parsed_sample):
    invalid = replace(
        parsed_sample,
        statement_period_from=parsed_sample.statement_period_to,
        statement_period_to=parsed_sample.statement_period_from,
    )
    staged = stage_parsed_shopee_weekly_statement(invalid)

    assert staged.result == NEEDS_REVIEW
    assert "invalid_statement_period" in {
        issue.code for issue in staged.validation_issues
    }
    assert staged.order_reconciliations == ()


def test_reconciliation_uses_final_estimated_and_unmatched_rules(parsed_sample):
    first, second, third, fourth = parsed_sample.order_rows[:4]
    existing_orders = [
        {
            "platform": "Shopee",
            "order_id": first.order_id,
            "income_type": "Final",
            "order_income": first.total_released_amount,
        },
        {
            "platform": "Shopee",
            "order_id": second.order_id,
            "income_type": "Final",
            "order_income": second.total_released_amount + Decimal("1.00"),
        },
        {
            "platform": "Shopee",
            "order_id": third.order_id,
            "income_type": "Estimated",
            "order_income": third.total_released_amount + Decimal("2.00"),
        },
    ]
    staged = stage_parsed_shopee_weekly_statement(
        parsed_sample, existing_orders=existing_orders
    )
    by_order = {item.order_id: item for item in staged.order_reconciliations}

    assert by_order[first.order_id].status == "Matched"
    assert by_order[second.order_id].status == "Different"
    assert by_order[third.order_id].status == "Estimated Only"
    assert by_order[fourth.order_id].status == "Unmatched Order"
    assert staged.result == NEEDS_REVIEW


def test_reconciliation_uses_rm002_tolerance_without_blocking_ready_state(
    parsed_sample,
):
    within_tolerance, beyond_tolerance = parsed_sample.order_rows[:2]
    staged = stage_parsed_shopee_weekly_statement(
        parsed_sample,
        existing_orders=[
            {
                "platform": "Shopee",
                "order_id": within_tolerance.order_id,
                "income_type": "Final",
                "order_income": within_tolerance.total_released_amount
                - Decimal("0.02"),
            },
            {
                "platform": "Shopee",
                "order_id": beyond_tolerance.order_id,
                "income_type": "Final",
                "order_income": beyond_tolerance.total_released_amount
                - Decimal("0.03"),
            },
        ],
    )
    by_order = {item.order_id: item for item in staged.order_reconciliations}

    assert by_order[within_tolerance.order_id].difference == Decimal("0.02")
    assert by_order[within_tolerance.order_id].status == "Matched"
    assert by_order[beyond_tolerance.order_id].difference == Decimal("0.03")
    assert by_order[beyond_tolerance.order_id].status == "Different"
    assert staged.result == NEEDS_REVIEW
    assert staged.validation_issues == ()
    assert "Underpayment" not in {
        item.status for item in staged.order_reconciliations
    }


def test_reconciliation_uses_shopee_platform_and_order_id_identity(parsed_sample):
    order_row = parsed_sample.order_rows[0]
    staged = stage_parsed_shopee_weekly_statement(
        parsed_sample,
        existing_orders=[
            {
                "platform": "Lazada",
                "order_id": order_row.order_id,
                "income_type": "Final",
                "order_income": order_row.total_released_amount,
            }
        ],
    )
    by_order = {item.order_id: item for item in staged.order_reconciliations}

    assert by_order[order_row.order_id].status == "Unmatched Order"
    assert staged.result == NEEDS_REVIEW


def test_existing_order_without_amount_is_not_reported_as_missing_order(parsed_sample):
    order_row = parsed_sample.order_rows[0]
    staged = stage_parsed_shopee_weekly_statement(
        parsed_sample,
        existing_orders=[
            {
                "platform": "Shopee",
                "order_id": order_row.order_id,
                "income_type": "Estimated",
                "order_income": None,
                "final_amount": None,
            }
        ],
    )
    by_order = {item.order_id: item for item in staged.order_reconciliations}

    assert by_order[order_row.order_id].status == "Missing Comparison Evidence"
    assert any("comparison evidence" in reason for reason in staged.review_reasons)


def _workbook_with_empty_income(source_bytes: bytes) -> bytes:
    output = BytesIO()
    empty_income = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<worksheet xmlns="http://schemas.openxmlformats.org/'
        b'spreadsheetml/2006/main"><dimension ref="A1"/><sheetData/></worksheet>'
    )
    with ZipFile(BytesIO(source_bytes)) as source, ZipFile(
        output, "w", ZIP_DEFLATED
    ) as destination:
        for item in source.infolist():
            if item.filename == "xl/worksheets/sheet5.xml":
                destination.writestr(item, empty_income)
            else:
                destination.writestr(item, source.read(item.filename))
    return output.getvalue()


def test_fallback_without_required_populated_income_is_rejected():
    empty_income = _workbook_with_empty_income(SAMPLE.read_bytes())
    staged = stage_shopee_weekly_statement(
        empty_income, source_filename="native-empty-income.xlsx"
    )

    assert staged.result == REJECTED
    assert staged.statement is None
    assert "Income is missing required column" in staged.rejection_reasons[0]


def test_corrupt_workbook_is_rejected():
    staged = stage_shopee_weekly_statement(
        b"not an xlsx workbook", source_filename="corrupt.xlsx"
    )

    assert staged.result == REJECTED
    assert staged.statement is None
    assert staged.rejection_reasons


def test_duplicate_gates_do_not_create_a_second_import(parsed_sample):
    exact = StatementReference(
        file_hash=parsed_sample.file_hash,
        statement_period_from=parsed_sample.statement_period_from,
        statement_period_to=parsed_sample.statement_period_to,
    )
    exact_result = stage_parsed_shopee_weekly_statement(
        parsed_sample, existing_statements=[exact]
    )

    assert exact_result.result == ALREADY_IMPORTED
    assert exact_result.duplicate_status == "ALREADY_IMPORTED"
    assert exact_result.already_imported is True
    assert exact_result.validation_issues == ()
    assert exact_result.review_reasons == ()
    assert exact_result.order_reconciliations == ()
    assert exact_result.adjustment_reconciliations == ()
    assert exact_result.eligible_for_future_atomic_commit is False

    revised = replace(exact, file_hash="different-file-hash")
    revised_result = stage_parsed_shopee_weekly_statement(
        parsed_sample, existing_statements=[revised]
    )

    assert revised_result.result == NEEDS_REVIEW
    assert revised_result.duplicate_status == "POSSIBLE_REVISION"
    assert revised_result.already_imported is False
    assert revised_result.eligible_for_future_atomic_commit is False


def test_adjustment_linking_is_non_blocking(parsed_sample):
    linked = parsed_sample.adjustments[0].linked_order_id
    staged = stage_parsed_shopee_weekly_statement(
        parsed_sample,
        existing_orders=[
            {
                "platform": "Shopee",
                "order_id": linked,
                "income_type": "Final",
                "order_income": "0.00",
            }
        ],
    )
    counts = Counter(item.status for item in staged.adjustment_reconciliations)

    assert counts == {"Matched": 1, "Unmatched Adjustment": 2}
    assert staged.result == NEEDS_REVIEW
