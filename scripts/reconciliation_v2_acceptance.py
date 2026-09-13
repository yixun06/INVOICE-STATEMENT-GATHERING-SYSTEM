"""Local-only, read-only acceptance harness for Shopee Reconciliation V2."""

from __future__ import annotations

import argparse
from collections import Counter
from decimal import Decimal
import json
from pathlib import Path
import sys

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from src.invoice_app.domain.historical_invoice import InvoiceBundle
from src.invoice_app.domain.statement_reconciliation_v2 import IdentityScope
from src.invoice_app.parsers.shopee_weekly_statement_parser import (
    parse_shopee_weekly_statement,
)
from src.invoice_app.repositories.google_sheets_historical_invoice_repository import (
    _deserialize_snapshot,
)
from src.invoice_app.services.product_price_master import ProductPriceMaster
from src.invoice_app.services.shopee_statement_item_matching import (
    product_family_resolver_from_price_master,
)
from src.invoice_app.services.shopee_statement_reconciliation_v2 import (
    evaluate_statement_reconciliation,
)


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Read local workbook snapshots and print aggregate Reconciliation V2 "
            "acceptance evidence. Inputs are never modified."
        )
    )
    parser.add_argument("--uat2-xlsx", type=Path, required=True)
    parser.add_argument("--statement-xlsx", type=Path, required=True)
    parser.add_argument("--product-master-xlsx", type=Path, required=True)
    args = parser.parse_args()

    bundles = _read_invoice_bundles(args.uat2_xlsx)
    statement = parse_shopee_weekly_statement(args.statement_xlsx)
    product_master = ProductPriceMaster.from_xlsx(args.product_master_xlsx)
    result = evaluate_statement_reconciliation(
        statement,
        tuple(bundle.order for bundle in bundles),
        tuple(item for bundle in bundles for item in bundle.items),
        product_families=product_family_resolver_from_price_master(product_master),
    )

    invoice_order_ids = {bundle.order.order_id for bundle in bundles}
    statement_order_ids = {row.order_id for row in statement.order_rows}
    invoice_items = tuple(item for bundle in bundles for item in bundle.items)
    member_scopes = Counter(
        identity.identity_scope.value
        for order_result in result.orders
        for identity in order_result.evidence.identities
        for _ in identity.statement_members
    )
    settlement = Counter(
        order_result.summary.settlement_basis.value for order_result in result.orders
    )
    invoice_merchandise = sum(
        (bundle.order.product_price or Decimal("0") for bundle in bundles),
        Decimal("0"),
    )
    statement_merchandise = sum(
        (
            row.financial_components.get("Product Price") or Decimal("0")
            for row in statement.order_rows
        ),
        Decimal("0"),
    )
    unexplained_residual = sum(
        (
            order_result.evidence.settlement.unexplained_residual or Decimal("0")
            for order_result in result.orders
            if order_result.summary.settlement_basis.value == "NONE"
        ),
        Decimal("0"),
    )
    report = {
        "order_population": {
            "invoice": len(invoice_order_ids),
            "statement": len(statement_order_ids),
            "bidirectional_overlap": len(invoice_order_ids & statement_order_ids),
            "invoice_only": len(invoice_order_ids - statement_order_ids),
            "statement_only": len(statement_order_ids - invoice_order_ids),
        },
        "product_population": {
            "invoice_items": len(invoice_items),
            "statement_sku_rows": len(statement.sku_rows),
            "statement_members_classified": sum(member_scopes.values()),
        },
        "identity_member_counts": {
            scope.value: member_scopes[scope.value] for scope in IdentityScope
        },
        "coverage": {
            "duplicate_item_consumption": len(
                result.coverage.duplicate_invoice_consumption
            ),
            "uncovered_statement_members": len(
                result.coverage.uncovered_statement_members
            ),
            "uncovered_invoice_members": len(result.coverage.uncovered_invoice_members),
        },
        "merchandise": {
            "invoice_product_price": _money(invoice_merchandise),
            "statement_product_price": _money(statement_merchandise),
            "residual": _money(statement_merchandise - invoice_merchandise),
            "reconciled_orders": sum(
                order_result.summary.merchandise_reconciled
                for order_result in result.orders
            ),
            "total_orders": len(result.orders),
        },
        "settlement": {
            "EXACT": settlement["EXACT"],
            "EXPLAINED": settlement["EXPLAINED"],
            "NONE": settlement["NONE"],
            "sum_unexplained_residual": _money(unexplained_residual),
        },
        "adjustment_total_separate": _money(result.statement_adjustment_total),
        "quantity": {
            "invoice_items_available": sum(
                item.quantity is not None for item in invoice_items
            ),
            "invoice_items_total": len(invoice_items),
            "statement_quantity_available": False,
        },
        "statement_validation_issue_count": len(result.statement_validation_issues),
        "product_family_snapshot_sha256": result.product_family_snapshot.sha256,
        "rule_version": result.rule_version,
    }
    print(json.dumps(report, indent=2, sort_keys=True))


def _read_invoice_bundles(source: Path) -> tuple[InvoiceBundle, ...]:
    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        tabs = {
            name: tuple(
                tuple(cell.value for cell in row)
                for row in workbook[name].iter_rows()
            )
            for name in ("Invoice_Orders", "Invoice_Items")
        }
    finally:
        workbook.close()
    return tuple(_deserialize_snapshot(tabs, 0).bundles.values())


def _money(value: Decimal) -> str:
    return format(value.quantize(Decimal("0.01")), "f")


if __name__ == "__main__":
    main()
